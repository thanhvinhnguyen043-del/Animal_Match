from pathlib import Path
import hashlib, json, re
from PIL import Image
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
WORK = ROOT / 'bento_report_evidence_audit'
WORK.mkdir(exist_ok=True)
REF = ROOT.parent / 'Báo cáo BentoLoop tối ưu nội dung.docx'
OUT = ROOT.parent / 'Báo cáo BentoLoop bổ sung dẫn chứng mục 3 4 5.docx'
TEMP = Path(r'C:\Users\ADMIN\AppData\Local\Temp')
source_hash = hashlib.sha256(REF.read_bytes()).hexdigest()
d = Document(REF)
original = (ROOT / 'build_bento_report.py').read_text(encoding='utf-8-sig')
helpers = original[original.index('def para(text):'):original.index('d.add_page_break()')]
exec(compile(helpers, str(ROOT / 'build_bento_report.py'), 'exec'))
figures = [None] * 17

def get_para(prefix):
    return next(p for p in d.paragraphs if p.text.startswith(prefix) and not p.style.name.lower().startswith('toc'))

def image_file(folder, n, caption, width=16, box=None):
    path = ROOT / folder / f'figure-{n:02d}.png'
    with Image.open(path) as im:
        bounds = box or (0, 0, im.width, im.height)
    return fig(path, bounds, caption, width)

def raw(uid, box, caption, width=16):
    candidates = [ROOT / 'bento_report_expanded' / 'evidence' / f'codex-clipboard-{uid}.png',
                  ROOT / 'bento_report_final' / 'evidence' / f'codex-clipboard-{uid}.png',
                  TEMP / f'codex-clipboard-{uid}.png']
    source = next(p for p in candidates if p.exists())
    saved = WORK / 'evidence' / source.name
    saved.parent.mkdir(exist_ok=True)
    saved.write_bytes(source.read_bytes())
    return fig(saved, box, caption, width)

def insert_before(anchor, writer):
    baseline = set(d.element.body)
    writer()
    new_elements = [el for el in d.element.body if el not in baseline and el.tag != qn('w:sectPr')]
    for el in new_elements:
        anchor.addprevious(el)

# Section 3: retain the existing input, return, key, JSON and runtime evidence.
# Insert the missing connecting branch rather than repeating the same code.
anchor = get_para('Hình 4.')._p.getprevious()
def branch_evidence():
    n = raw('b90788c8-59bd-4dba-a17e-6a7a2a451e6c', (1068, 341, 1882, 558),
            'Các dòng 27 đến 38 nối kết quả Equals với các nhánh kiểm tra bổ sung', 16.8)
    para(f'Hình {n} bổ sung phần giữa hàm. Điều kiện ở dòng 27 mở nhánh khi Equals không khớp. Dòng 31 gọi CheckTransformedLicense; nếu kết quả vẫn bằng 0, dòng 33 gọi DispatchVerification. Dòng 34 chuyển tới LAB_01f76914 khi lần kiểm tra này cũng thất bại. Ngược lại, khi Equals đã khớp, khối kiểm tra bổ sung được bỏ qua và luồng tiếp tục tới phần trả về thành công bên dưới.')
insert_before(anchor, branch_evidence)
p = get_para('Hình 4 cho thấy')
p.text = 'Hình 4 cho thấy dòng 51 gán uVar3 = 1 và dòng 57 trả về giá trị đó. Tại LAB_01f76914, dòng 55 gán uVar3 = 0. Ghép phần giữa hàm với phần cuối xác định được đường trả về của hai trường hợp, từ đó chọn StringLiteral_2770 để đọc chuỗi được nhánh trực tiếp chấp nhận.'

anchor = get_para('0x043FCE70 +')._p
def base_evidence():
    n = image_file('bento_report', 7, 'Memory Map hiển thị Image Base 00100000', 12.2, (0, 0, 450, 78))
    para(f'Thanh tiêu đề Memory Map trong Hình {n} ghi Image Base: 00100000, tương ứng 0x100000. Đây là giá trị dùng trong phép quy đổi địa chỉ JSON sang địa chỉ Listing, cũng là giá trị đã dùng để mở VerifyLicense ở mục 2.2.')
insert_before(anchor, base_evidence)

# Section 4: export settings existed; add the actual export result and the file
# bearing its final name. Do not add a second patch-operation screenshot.
anchor = get_para('4.3 Thay thư viện')._p
def export_evidence():
    n = image_file('bento_report_expanded', 19, 'Export Results Summary ghi file đích và định dạng Original File', 15.5)
    para(f'Hình {n} là kết quả sau khi xuất: Destination file trỏ tới C:\\Thuchanh2\\libil2cpp_patched và Format là Original File. Ảnh này kiểm chứng bước Export đã tạo file đích, thay vì chỉ lưu thay đổi trong project Ghidra.')
    n = raw('d2d58be5-05c5-4dc3-9ea4-af577e6a4b7e', (176, 247, 775, 313),
            'File thư viện mang tên libil2cpp.so trước khi copy', 15.0)
    para(f'Hình {n} ghi lại file dưới tên libil2cpp.so trước khi copy. Tên này khớp thư viện trong lib\\arm64-v8a; hộp thoại thay file ở bước tiếp theo ghi rõ cùng tên và đúng thư mục đích.')
insert_before(anchor, export_evidence)

# Section 5: replace the unsupported single troubleshooting paragraph with
# concise cause -> evidence -> action pairs. Preserve output and runtime proof.
old = get_para('Quá trình build cần cập nhật')
anchor = old._p.getnext()
d.element.body.remove(old._p)
def build_evidence():
    heading('5.1.1 Lỗi manifest khi dùng Apktool cũ', 3)
    para('Lần Compile với Apktool 2.6.1 dừng ở manifest. Log dưới đây chỉ ra lỗi unexpected element <property>, vì công cụ build chưa xử lý được phần tử này trong APK.')
    n = raw('e8b9bc9e-248f-47eb-a19b-44ddec9530bf', (370, 389, 1035, 569),
            'Log build báo unexpected element property trong AndroidManifest.xml', 16.8)
    para(f'Các dòng lỗi trong Hình {n} đều trỏ tới AndroidManifest.xml và phần tử property. Bước xử lý là thêm Apktool 2.12.1 vào APK Easy Tool rồi chọn phiên bản mới để Compile lại. Lỗi này thuộc khâu build tài nguyên, chưa phải kết quả chạy VerifyLicense đã sửa.')
    heading('5.1.2 Bỏ tham số không còn được hỗ trợ', 3)
    n = raw('909097fe-4291-45f5-91f5-60a9e0b690b1', (484, 530, 1146, 635),
            'Apktool 2.12.1 báo Unrecognized option với use aapt2', 16.8)
    para(f'Hình {n} cho thấy lệnh build đã gọi apktool_2.12.1.jar nhưng vẫn kèm --use-aapt2, sau đó báo Unrecognized option. Cần chỉnh tùy chọn trong giao diện APK Easy Tool để bỏ tham số cũ này.')
    n = image_file('bento_report_final', 18, 'Bỏ tích Use the aapt2 binary instead of aapt trong phần Compile', 15.6, (8, 364, 662, 450))
    para(f'Trong Hình {n}, ô Use the aapt2 binary instead of aapt không được tích. Apktool mới dùng AAPT2 mặc định, nên bỏ ô này giúp giao diện ngừng thêm --use-aapt2 vào lệnh build. Các tùy chọn khác không cần thay để xử lý lỗi đang gặp.')
    heading('5.1.3 Làm mới framework trước khi Compile lại', 3)
    n = image_file('bento_report_final', 19, 'Build tiếp tục nhưng báo thiếu enableOnBackInvokedCallback', 16.8)
    para(f'Hình {n} hiển thị Using Apktool 2.12.1 và Building resources with aapt2, xác nhận quá trình đã qua lỗi tham số. Các dòng tiếp theo báo attribute android:enableOnBackInvokedCallback not found. Framework đang được lưu trong bộ nhớ đệm cần được làm mới để nhận thuộc tính này.')
    n = image_file('bento_report_final', 20, 'Framework mới 1.apk nằm cạnh bản sao cũ 1.apk.bak', 15.4)
    para(f'Đổi tên framework\\1.apk trong bộ nhớ đệm Apktool thành 1.apk.bak rồi Compile lại. Hình {n} cho thấy file 1.apk mới và bản sao 1.apk.bak cùng nằm trong thư mục framework, có thời điểm sửa khác nhau. Sau bước này, công cụ build và ký được APK như ảnh kết quả bên dưới.')
    heading('5.1.4 Kết quả build và ký APK', 3)
insert_before(anchor, build_evidence)

p = get_para('Hình 11 cho thấy')
p.text = p.text.replace('Hai lệnh cấp phát stack ban đầu', 'Hai lệnh chuẩn bị stack ban đầu')

# Give readable figure sizes to the code excerpts by removing empty right-hand
# space; the exact source pixels and relevant line numbers remain intact.
replacements = {
    'Hình 3.': ('bento_report_expanded', 8, (0, 0, 690, 183), 16.3),
    'Hình 4.': ('bento_report_expanded', 9, (0, 0, 690, 371), 14.5),
}
for prefix, (folder, n, box, width) in replacements.items():
    cap = get_para(prefix)
    picture = cap._p.getprevious()
    source = ROOT / folder / f'figure-{n:02d}.png'
    crop_path = WORK / f'readable-{n:02d}.png'
    Image.open(source).crop(box).save(crop_path)
    from docx.text.paragraph import Paragraph
    p = Paragraph(picture, d)
    p.clear()
    p.add_run().add_picture(str(crop_path), width=Cm(width))

# Renumber all figure references in a single pass from the original IDs to their
# final reading order, including the captions of newly inserted evidence.
mapping = {}
for p in d.paragraphs:
    match = re.match(r'^Hình (\d+)\.', p.text)
    if match:
        mapping[int(match.group(1))] = len(mapping) + 1
for p in d.paragraphs:
    if re.search(r'Hình \d+', p.text):
        new_text = re.sub(r'Hình (\d+)', lambda m: f'Hình {mapping[int(m.group(1))]}', p.text)
        if p.text != new_text:
            p.text = new_text
        if re.match(r'^Hình \d+\.', p.text):
            for r in p.runs:
                r.bold = True
                r.italic = True
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(0, 0, 0)
used = {el.get(qn('r:embed')) for el in d.element.xpath('.//a:blip')}
for rid, rel in list(d.part.rels.items()):
    if rel.reltype.endswith('/image') and rid not in used:
        d.part.drop_rel(rid)
d.save(OUT)
assert hashlib.sha256(REF.read_bytes()).hexdigest() == source_hash
check = Document(OUT)
text = '\n'.join(check.element.xpath('.//w:t/text()'))
assert 'SDK' not in text
assert not re.search(r'\b(tôi|chúng tôi|chúng ta|mình|bạn)\b', text, re.I)
assert len(mapping) == 26 and len(check.inline_shapes) == 27
imgs = [r.target_part.blob for r in check.part.rels.values() if r.reltype.endswith('/image')]
assert len(imgs) == len(set(hashlib.sha256(b).digest() for b in imgs))
manifest = {'reference': str(REF), 'reference_sha256': source_hash, 'output': str(OUT),
            'added_figures': figures[17:], 'figure_mapping': mapping, 'figure_count': 26,
            'gaps_closed': ['middle branch of VerifyLicense', 'Image Base supporting address conversion', 'export result', 'renamed library file', 'property build failure', 'unsupported option', 'option disabled', 'framework attribute failure', 'framework refreshed'],
            'retained_sufficient_evidence': ['key and JSON', 'original key test and Level 1', 'native instructions before and after patch', 'Decompiler return 1', 'export settings', 'library replacement', 'signed APK and output file', 'rebuilt APK accepts vinh123'],
            'render_pending': True}
(WORK / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
print('OUTPUT', OUT)
print('FIGURES', len(mapping), 'ADDED', len(figures) - 17)
