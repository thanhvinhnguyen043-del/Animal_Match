from pathlib import Path
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
import hashlib,json,re
from PIL import Image
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT=Path(__file__).resolve().parent
WORK=ROOT/'bento_report';WORK.mkdir(exist_ok=True)
REF=ROOT.parent/'Báo cáo dịch ngược SandSort rà soát ảnh và nội dung.docx'
OUT=ROOT.parent/'Báo cáo dịch ngược BentoLoop tìm license key.docx'
TEMP=Path(r'C:\Users\ADMIN\AppData\Local\Temp')
ref_hash=hashlib.sha256(REF.read_bytes()).hexdigest()
d=Document(REF)
cover=d.tables[0]._tbl
first=d.element.body[0]
for el in list(d.element.body):
    if el is not first and el is not cover and el.tag!=qn('w:sectPr'):d.element.body.remove(el)

# Reuse the academic cover and its existing author information.
for el in cover.xpath('.//w:p'):
    p=Paragraph(el,d)
    if p.text.strip()=='DỊCH NGƯỢC':p.style=d.styles['Title']
    if p.text.strip()=='AN TOÀN GAME':
        slot=p._p.getnext()
        while slot is not None and slot.tag!=qn('w:p'):slot=slot.getnext()
        assert slot is not None
        sp=Paragraph(slot,d)
        assert not sp.text.strip()
        sp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=sp.add_run('BENTOLOOP');r.bold=True;r.font.name='Times New Roman';r.font.size=Pt(17)
        sp.paragraph_format.space_before=Pt(0);sp.paragraph_format.space_after=Pt(0)
        break

normal=d.styles['Normal']
normal.font.name='Arial';normal.font.size=Pt(11);normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.space_before=Pt(0);normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.first_line_indent=Cm(0.8)
for name,size in [('Heading 1',12),('Heading 2',11.5),('Heading 3',11)]:
    s=d.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.first_line_indent=Pt(0);s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(5)
    s.paragraph_format.keep_with_next=True;s.paragraph_format.page_break_before=False

def para(text):
    p=d.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return p
def heading(text,level=1):
    p=d.add_paragraph(text,style=f'Heading {level}');p.paragraph_format.page_break_before=False
    return p
def code(text):
    p=d.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_together=True
    p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(6)
    r=p.add_run(text);r.font.name='Consolas';r.font.size=Pt(9.5)
    return p
figures=[]
def fig(source,box,caption,width=17):
    if not isinstance(source,Path):source=TEMP/f'codex-clipboard-{source}.png'
    im=Image.open(source).crop(box)
    n=len(figures)+1;dest=WORK/f'figure-{n:02d}.png';im.save(dest)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pf=p.paragraph_format;pf.first_line_indent=Pt(0);pf.left_indent=Pt(0);pf.right_indent=Pt(0)
    pf.line_spacing=1;pf.space_before=Pt(3);pf.space_after=Pt(3);pf.keep_with_next=True
    # Keep long screenshots readable while fitting within a page with their caption.
    max_height=15.3
    width=min(width,max_height*im.width/im.height)
    picture=p.add_run().add_picture(str(dest),width=Cm(width))
    picture._inline.docPr.set('descr',caption)
    cp=d.add_paragraph(f'Hình {n}. {caption}');cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.first_line_indent=Pt(0);cp.paragraph_format.line_spacing=1
    cp.paragraph_format.space_before=Pt(0);cp.paragraph_format.space_after=Pt(6)
    cp.paragraph_format.keep_together=True;cp.paragraph_format.keep_with_next=False
    for r in cp.runs:r.bold=True;r.italic=True;r.font.size=Pt(10.5);r.font.color.rgb=RGBColor(0,0,0)
    figures.append({'number':n,'source':str(source),'crop':box,'caption':caption,'file':str(dest)})
    return n
def table(headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,(cell,label,w) in enumerate(zip(t.rows[0].cells,headers,widths)):
        cell.width=Cm(w);cell.text=label
    for values in rows:
        cells=t.add_row().cells
        for cell,value,w in zip(cells,values,widths):cell.width=Cm(w);cell.text=value
    for gc,w in zip(t._tbl.tblGrid.gridCol_lst,widths):gc.set(qn('w:w'),str(int(Cm(w).twips)))
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            cell._tc.get_or_add_tcPr().append(margins)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');cell._tc.get_or_add_tcPr().append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.05
                p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
                for r in p.runs:r.font.size=Pt(10.5);r.bold=ri==0
        if ri==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

d.add_page_break()
p=d.add_paragraph('MỤC LỤC');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
for r in p.runs:r.bold=True;r.font.size=Pt(18)
p=d.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
r=p.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');r._r.append(begin)
r=p.add_run();instruction=OxmlElement('w:instrText');instruction.set(qn('xml:space'),'preserve');instruction.text=' TOC \\o "1-2" \\h \\z \\u ';r._r.append(instruction)
r=p.add_run();sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
p.add_run('Mục lục được cập nhật khi mở tài liệu.')
r=p.add_run();end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.append(end)
d.add_page_break()

heading('1 Mục tiêu và môi trường thực hành')
para('Bài thực hành phân tích cơ chế kiểm tra license của BentoLoop trên Android. Từ dữ liệu dump và mã native, đã tìm được chuỗi BENTO-EASY-2025 ở nhánh so sánh trực tiếp của VerifyLicense. Chuỗi này đã được nhập vào APK gốc để vượt màn hình license và vào Level 1. Phần sửa thư viện và build lại APK chưa được thực hiện.')
para('Theo yêu cầu thực hành, mỗi bước cần nêu dữ liệu đầu vào, thao tác và kết quả kèm ảnh. Hai kết quả cần đạt là tìm key để vượt màn hình kiểm tra và sửa hoặc vòng tránh cơ chế kiểm tra rồi build lại APK. Báo cáo trình bày phần tìm key đã hoàn thành và phân biệt với nhiệm vụ sửa APK còn lại.')
table(['Thành phần','Thông tin sử dụng'],[
 ('Môi trường phân tích','Windows trong máy ảo VMware Workstation, tên thuchanhdichnguoc1'),
 ('APK','BentoLoop.apk; package com.ig.bento.loop; phiên bản 1.0.0, version code 5'),
 ('APK Easy Tool','Phiên bản 1.60, Apktool 2.6.1; dùng để Decompile APK'),
 ('Dữ liệu IL2CPP','Gói libil2cpp_il2cppdump_20261007_0259 do giảng viên cung cấp'),
 ('dnSpy và dump.cs','Tra cứu cấu trúc lớp, chữ ký hàm và địa chỉ RVA'),
 ('Ghidra','Phiên bản 12.1.2; project Thuchanh2; thư viện ARM64'),
 ('Kiểm chứng','Chạy Bento Loop trên Android, nhập key tìm được và kiểm tra vào Level 1'),
],[4.1,12.9])

heading('2 Decompile APK và xác định thư viện')
para('Đầu vào là C:\\Thuchanh2\\BentoLoop.apk. Trong APK Easy Tool, Browse được dùng để chọn APK, sau đó bấm Decompile. Kết quả hiển thị Decompile successful. Công cụ cũng hiển thị tên package, phiên bản và thông tin SDK của ứng dụng.')
fig('74a36e31-fa57-40c9-8e14-cb0f5c2a950c',(620,193,1312,795),'APK Easy Tool chọn BentoLoop.apk và báo Decompile successful',width=15)
para('Đầu ra nằm trong C:\\Users\\DucNM\\Documents\\APK Easy Tool\\1-Decompiled APKs\\BentoLoop. Thư mục chứa assets, lib, res, các thư mục smali, AndroidManifest.xml và apktool.yml. Đây là thư mục làm việc để lấy thư viện native và dùng khi build lại APK ở nhiệm vụ tiếp theo.')
fig('b6891c8b-bf62-425a-bc4c-b8c8368b8a5b',(307,144,1002,612),'Thư mục BentoLoop sau Decompile có lib, assets và apktool.yml',width=15)
para('Trong lib\\arm64-v8a, libil2cpp.so có dung lượng hiển thị 71 122 KB. File này được chọn để phân tích thay cho các thư viện quảng cáo, Firebase hoặc thư viện hỗ trợ khác trong cùng thư mục. Đường dẫn đầy đủ là C:\\Users\\DucNM\\Documents\\APK Easy Tool\\1-Decompiled APKs\\BentoLoop\\lib\\arm64-v8a\\libil2cpp.so.')
fig('e461671f-2ee1-4d24-8c97-602bf2eb0fa9',(410,220,1086,643),'libil2cpp.so trong thư mục ARM64 của BentoLoop',width=16)

heading('3 Tra cứu lớp LicenseVerifier trong dữ liệu dump')
para('Đầu vào của bước tra cứu là gói dump có sẵn trong C:\\Thuchanh2\\libil2cpp_il2cppdump_20261007_0259. Các file được sử dụng gồm dump.cs, script.json, stringliteral.json, ghidra.py và thư mục DummyDll. Thư viện dùng trong Ghidra được lấy từ APK vừa Decompile. Không thực hiện lại bước tạo dump trong phần này.')
para('Tìm LicenseVerifier trong dump.cs cho kết quả lớp static thuộc namespace GeckoArrow.Runtime.LoadingSystem. Lớp có field private static readonly byte[] EncryptedBytes. VerifyLicense nhận string inputKey và trả về bool; CheckTransformedLicense nhận string key và cũng trả về bool. Hai hàm này là đầu mối để đọc cách kiểm tra license trong mã native.')
fig('60e41892-2a80-49fc-8ed1-4b1a222784f4',(203,453,867,869),'Khai báo LicenseVerifier, mảng EncryptedBytes và RVA của các phương thức trong dump.cs',width=16)
table(['Phương thức','RVA','Địa chỉ Ghidra'],[
 ('VerifyLicense','0x1E7689C','0x1F7689C'),
 ('CheckTransformedLicense','0x1E76FA0','0x1F76FA0'),
 ('IsLicenseActivated','0x1E76EF4','0x1F76EF4'),
 ('SaveLicenseActivated','0x1E769FC','0x1F769FC'),
 ('.cctor','0x1E77110','0x1F77110'),
],[7.5,4.75,4.75])
para('Các thân hàm rỗng trong dump.cs chỉ phục vụ tra cứu. Kết quả bước này là xác định đúng tên lớp, vai trò các phương thức và RVA để mở mã native. Địa chỉ Ghidra trong bảng được tính sau khi kiểm tra Image Base ở bước tiếp theo.')

heading('4 Import thư viện và gắn tên trong Ghidra')
heading('4.1 Nhận dạng thư viện ARM64',2)
para('Đầu vào là libil2cpp.so trong thư mục arm64-v8a của BentoLoop. File được import vào project Thuchanh2. Ghidra nhận dạng định dạng Executable and Linking Format và language AARCH64:LE:64:v8A:default. Sau khi mở CodeBrowser, bước phân tích tự động được chạy để nhận diện lệnh, hàm và tham chiếu.')
fig('479baaa6-3307-424d-bd00-0b38434dae47',(580,391,1080,642),'Hộp Import nhận dạng libil2cpp.so là ELF ARM64',width=14)
para('Import Results Summary xác nhận processor AARCH64, little-endian, địa chỉ 64 bit và ELF shared object. MD5 của thư viện là 36834a117aa0a39f3e2a1068a5c6e846. Mã nhận diện này được dùng để phân biệt đúng file đã phân tích, không phải mã băm của toàn bộ APK.')
fig('3aaaa8b0-b063-4df0-b106-173ffb6fd224',(476,155,1214,749),'Thông tin kiến trúc và mã nhận diện thư viện trong Import Results Summary',width=15)
heading('4.2 Quy đổi RVA sang địa chỉ Ghidra',2)
para('Memory Map hiển thị Image Base 00100000, tương ứng 0x100000. Vì vậy, RVA trong dump cần cộng thêm Image Base để mở đúng vị trí trong chương trình đã import. Với VerifyLicense, kết quả là 0x1F7689C. File Offset là vị trí trong file, không dùng thay cho địa chỉ Ghidra.')
code('Địa chỉ Ghidra = Image Base + RVA\n0x100000 + 0x1E7689C = 0x1F7689C')
fig('8e864c30-d6f7-483b-9b6d-7dd7ab5f3a5f',(207,164,1121,483),'Memory Map xác nhận Image Base 0x100000 và các vùng bộ nhớ',width=17)
heading('4.3 Áp dụng tên từ script.json',2)
para('Thư mục chứa ghidra.py được thêm vào Script Manager. Khi chạy script, script.json trong gói dump Bento được chọn làm đầu vào. Dữ liệu này gắn tên các phương thức và chuỗi vào địa chỉ tương ứng. Đầu ra thể hiện ở Decompiler với tên đầy đủ như GeckoArrow.Runtime.LoadingSystem.LicenseVerifier$$VerifyLicense và các nhãn StringLiteral.')
fig('b4fe4676-83d6-49cd-b7d2-996c1fbd8940',(670,230,1452,825),'Chọn script.json của Bento làm đầu vào cho script gắn tên',width=15)
para('Trước khi gắn tên, Ghidra dùng tên tự động FUN_01f7689c. Sau khi gắn tên, cùng địa chỉ đó hiển thị VerifyLicense. Việc thay đổi nhãn làm mã giả dễ đọc hơn, không thay đổi mã máy trong thư viện. File gamebento.py được cung cấp cũng là script nhập dữ liệu script.json; chức năng của nó là gắn nhãn, chú thích và tạo hàm, không giải mã key.')

heading('5 Phân tích luồng kiểm tra license')
heading('5.1 Kiểm tra đầu vào và nhánh so sánh trực tiếp',2)
para('Đầu vào của bước phân tích là hàm VerifyLicense tại 0x1F7689C. Hàm được mở bằng Go To, sau đó đọc trong Decompiler. Phần đầu gọi String.IsNullOrWhiteSpace để loại đầu vào rỗng hoặc chỉ có khoảng trắng. Với đầu vào hợp lệ, String.Trim tạo chuỗi đã bỏ khoảng trắng ở hai đầu.')
para('Tiếp theo, String.Equals so sánh chuỗi đã Trim với StringLiteral_2770, truyền giá trị 4 ở tham số kiểu so sánh. Khi kết quả khác 0, luồng đi đến nhánh ngoài phần kiểm tra bổ sung. Khi kết quả bằng 0, VerifyLicense gọi CheckTransformedLicense, rồi có thể gọi SecurityAuthService.DispatchVerification nếu nhánh trước chưa chấp nhận đầu vào.')
fig('94d8ddd2-1581-4280-bd08-3825a4ab5540',(1075,496,1916,829),'VerifyLicense chuẩn hóa đầu vào rồi so sánh với StringLiteral_2770',width=17)
heading('5.2 Nhánh trả về thành công',2)
para('Trong phần cuối hàm, nhánh chuỗi khớp với StringLiteral_2770 chọn StringLiteral_1120 để ghi log. Luồng sau đó gán giá trị trả về bằng 1. Đầu vào không được chấp nhận bởi các nhánh kiểm tra chuyển đến LAB_01f76914 và trả về 0. Như vậy, có thể tìm key ở chuỗi được so sánh trực tiếp trước khi giải nhánh mã hóa.')
fig('b90788c8-59bd-4dba-a17e-6a7a2a451e6c',(1075,568,1916,939),'Nhánh so sánh khớp dẫn đến giá trị trả về 1 trong VerifyLicense',width=17)
para('Luồng chính có thể viết lại bằng mã giả dưới đây. Đoạn này mô tả điều kiện đã quan sát, không phải mã nguồn C# gốc.')
code('if (IsNullOrWhiteSpace(inputKey)) return false;\nkey = Trim(inputKey);\nif (Equals(key, StringLiteral_2770, 4)) return true;\nif (CheckTransformedLicense(key)) return true;\nreturn SecurityAuthService.DispatchVerification(key);')
heading('5.3 Nhánh kiểm tra mảng mã hóa',2)
para('CheckTransformedLicense tại 0x1F76FA0 kiểm tra độ dài chuỗi bằng độ dài mảng dữ liệu, sau đó duyệt từng vị trí. Vòng lặp lấy một byte trong mảng và một ký tự bằng String.get_Chars. Phép đối chiếu sử dụng XOR với 0x53 và giá trị (i × 13) AND 0xFF, với i bắt đầu từ 0.')
fig('e62eadeb-1c61-416a-9a91-5299ff88bcf8',(1282,165,1917,771),'Vòng lặp trong CheckTransformedLicense đối chiếu ký tự bằng phép XOR',width=14)
code('ký_tự[i] = EncryptedBytes[i] XOR 0x53 XOR ((i × 13) AND 0xFF)')
para('Hàm khởi tạo lớp .cctor tại 0x1F77110 tạo mảng byte dài 0x1E, tức 30 byte, rồi gọi RuntimeHelpers.InitializeArray để nạp dữ liệu từ field của PrivateImplementationDetails. Các byte tại nhãn tham chiếu field chưa phải nội dung mảng. Trong bài này, nhánh so sánh trực tiếp đã cung cấp key được game chấp nhận nên chưa thực hiện giải toàn bộ mảng 30 byte.')
fig('edcf5ad2-cb3f-42e6-8d71-6ef50936053e',(1103,165,1917,750),'Hàm .cctor tạo mảng 30 byte và gọi RuntimeHelpers.InitializeArray',width=16)

heading('6 Đối chiếu key và kiểm chứng trên APK gốc')
heading('6.1 Đọc nội dung StringLiteral_2770',2)
para('Đầu vào là nhãn StringLiteral_2770 trong lời gọi String.Equals của VerifyLicense. Bấm đúp nhãn trong Decompiler chuyển Listing đến 044fce70. Chú thích tại đây hiển thị BENTO-EASY-2025. Đây là chuỗi được dùng ở nhánh so sánh trực tiếp, không phải tên field EncryptedBytes hoặc khóa lưu trạng thái BENTO_LICENSE_ACTIVATED ở vị trí khác.')
fig('a6686564-c88d-4cf5-850e-f92e0706a53d',(598,484,1498,612),'StringLiteral_2770 tại 0x044FCE70 có nội dung BENTO-EASY-2025',width=17)
heading('6.2 Đối chiếu với stringliteral.json',2)
para('Trong stringliteral.json của cùng gói dump, tìm 43FCE70 cho mục có value là BENTO-EASY-2025 và address là 0x43FCE70. Cộng Image Base 0x100000 cho địa chỉ này được 0x44FCE70, khớp Listing. Cách hiển thị 044fce70 của Ghidra có thêm số 0 đầu và không ghi tiền tố 0x; đây vẫn là cùng một địa chỉ.')
code('0x043FCE70 + 0x00100000 = 0x044FCE70')
fig('8e3f048b-014f-43fa-88a2-d0d1bf1df6d4',(209,605,769,679),'Mục value và address của key trong stringliteral.json',width=12)
para('Chú thích Ghidra được script tạo từ dữ liệu dump nên việc mở JSON là đối chiếu nguồn của nhãn. Bằng chứng về vai trò của chuỗi nằm ở nhánh String.Equals và giá trị trả về trong VerifyLicense; kết quả chạy trên APK gốc xác nhận key có hiệu lực.')
heading('6.3 Nhập key và vượt màn hình license',2)
para('Đầu vào thử nghiệm là APK gốc chưa sửa và key BENTO-EASY-2025 đã xác định. Chuỗi được nhập vào ô license trên màn hình BENTO SECURITY SYSTEM, sau đó bấm XÁC NHẬN (VERIFY). Ảnh trước khi xác nhận thể hiện đúng key trong ô nhập.')
fig('9c379f73-aff8-4d3b-9b79-87f95c977b11',(680,337,1240,735),'Nhập BENTO-EASY-2025 vào màn hình kiểm tra license của APK gốc',width=12)
para('Sau khi xác nhận, màn hình license được vượt qua và game hiển thị Level 1. Kết quả này hoàn tất nhiệm vụ tìm một key được chấp nhận và kiểm chứng trên APK gốc. Kết quả chưa chứng minh rằng các chức năng khác đã được mở khóa hoặc mọi chuỗi nhập đều được chấp nhận.')
fig('40caf4ef-fb81-46c2-b684-283ae70bafe8',(680,40,1240,1032),'Bento Loop vào Level 1 sau khi xác nhận license',width=8)

heading('7 Trạng thái nhiệm vụ sửa và build APK')
para('Nhiệm vụ tìm key đã hoàn thành bằng nhánh so sánh trực tiếp và được kiểm chứng trên APK gốc. Chưa thực hiện patch hàm, thay libil2cpp.so, Compile, ZipAlign hoặc ký một APK đã sửa. Vì vậy, kết quả hiện tại chưa đáp ứng nhiệm vụ thứ hai là vòng tránh cơ chế kiểm tra và build lại APK.')
para('Theo hướng dẫn công cụ, bước build tiếp theo sử dụng thư mục BentoLoop đã Decompile và có apktool.yml. Sau khi có thư viện đã sửa, libil2cpp.so trong lib\\arm64-v8a được thay bằng bản đó, rồi chọn Select decompiled APK và Compile. APK kết quả cần được căn chỉnh và ký nếu công cụ chưa thực hiện tự động. Đây là quy trình cần thực hiện tiếp, chưa phải kết quả đã đạt.')

heading('8 Nhận xét an toàn và phương án phòng tránh')
heading('8.1 Rủi ro của nhánh so sánh key trên client',2)
para('BENTO-EASY-2025 được lưu dưới dạng chuỗi và được so sánh ngay trong thư viện phía client. Sau khi có dữ liệu dump và tên phương thức, có thể lần từ VerifyLicense đến chuỗi mà không cần giải nhánh EncryptedBytes. Việc đặt thêm một nhánh XOR không bảo vệ được key ở nhánh so sánh trực tiếp vẫn còn tồn tại.')
para('Một khả năng can thiệp khác là thay đổi kết quả trả về hoặc nhánh điều kiện trong hàm kiểm tra. Khả năng này chưa được thử nghiệm bằng APK đã patch trong bài hiện tại. Chưa đủ dữ liệu để kết luận game có cơ chế kiểm tra tính toàn vẹn hoặc xác thực bổ sung ngoài luồng đã xem hay không.')
heading('8.2 Phương án giảm rủi ro',2)
para('Với ứng dụng có máy chủ, license nên được xác minh trên máy chủ theo tài khoản hoặc quyền sử dụng. Client nhận kết quả hoặc token có thời hạn thay vì giữ một key chung để so sánh. Không đặt key bí mật hoặc nhánh chấp nhận thử nghiệm trong bản phát hành; các thao tác cần quyền sử dụng cũng phải được kiểm tra, thay vì chỉ khóa ở màn hình khởi động.')
para('Nếu cần hoạt động offline, có thể dùng dữ liệu license được ký số. Client giữ khóa công khai để kiểm tra chữ ký, còn khóa riêng dùng để phát hành license được giữ ngoài ứng dụng. Cách này tránh nhúng một bí mật chung để người dùng có thể đọc ra, nhưng mã kiểm tra phía client vẫn có thể bị patch nên cần kết hợp kiểm tra tính toàn vẹn và giới hạn phù hợp với môi trường offline.')

heading('9 Kết luận')
para('Qua bài thực hành, đã sử dụng APK Easy Tool để lấy thư viện native, đọc dump.cs để xác định lớp và RVA, rồi dùng Ghidra để đối chiếu địa chỉ và lần theo nhánh kiểm tra license. Kết hợp mã giả với stringliteral.json đã xác định được BENTO-EASY-2025. Thử nghiệm nhập key trên APK gốc cho kết quả vượt màn hình license và vào Level 1.')
para('Các bước đã thực hiện cho thấy có thể dùng bộ công cụ này để tìm điều kiện chấp nhận đầu vào, phân biệt tên nhãn với nội dung dữ liệu và kiểm chứng kết quả bằng việc chạy ứng dụng. Phần sửa thư viện để vòng tránh kiểm tra và build lại APK vẫn là nhiệm vụ cần thực hiện tiếp.')

# Remove old SandSort image relationships while retaining the cover logo.
used={el.get(qn('r:embed')) for el in d.element.xpath('.//a:blip')}
for rid,rel in list(d.part.rels.items()):
    if rel.reltype.endswith('/image') and rid not in used:d.part.drop_rel(rid)
d.core_properties.title='Báo cáo thực hành dịch ngược BentoLoop tìm license key'
d.core_properties.author='Nguyễn Thành Vinh'
d.core_properties.subject='Phân tích VerifyLicense và kiểm chứng BENTO-EASY-2025'
settings=d.settings.element
update=settings.find(qn('w:updateFields'))
if update is None:update=OxmlElement('w:updateFields');settings.append(update)
update.set(qn('w:val'),'true')
d.save(OUT)
assert hashlib.sha256(REF.read_bytes()).hexdigest()==ref_hash
check=Document(OUT)
caps=[p for p in check.paragraphs if re.match(r'^Hình \d+\.',p.text)]
assert len(caps)==16 and len(check.inline_shapes)==17
assert all(p.alignment==1 and all(r.bold and r.italic for r in p.runs if r.text) for p in caps)
assert not any('SandSort' in p.text for p in check.paragraphs)
manifest={'reference':str(REF),'reference_sha256':ref_hash,'output':str(OUT),'figures':figures,'key_test':'User confirmed original APK passes; two screenshots show entered key and Level 1','patch_build':'Not performed'}
(WORK/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
(WORK/'artifact.md').write_text(f'''# Bento report contract\nReference {REF}\nSHA256 {ref_hash}\nReference rendering previously inspected in report_work/compact_review/final_render, 30 pages.\nSingle portrait A4 section with source margins and footer preserved. Academic cover table and logo retained; an existing blank cover paragraph receives BentoLoop. All former SandSort body and figures removed. Body rebuilt using Arial 11, justified text, 1.12 line spacing, first-line indent 0.8 cm, after spacing 6 pt. Heading 1 Arial 12 bold, Heading 2 Arial 11.5 bold, black, no forced page breaks. Two initial page breaks separate cover, TOC and body. Figures centered, captions Arial 10.5 bold italic black; keep image and caption together. Tables use source-like gray headers and light gray grid with variable row height. Footer and identity information from reference preserved. Editable slots include all former body content, cover subject blank, document metadata and TOC. Runtime success confirmed by user and figures 15-16. No claimed patch or rebuilt APK. Retained source must remain unchanged.\n''',encoding='utf8')
print('REPORT_CREATED',OUT)
print('FIGURES',len(figures),'INLINE_SHAPES',len(check.inline_shapes))
