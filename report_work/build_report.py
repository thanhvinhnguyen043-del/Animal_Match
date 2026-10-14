from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

ROOT = Path(r'C:\Users\ADMIN\Downloads\My project (1)')
WORK = ROOT / 'report_work'
TEMP = Path(r'C:\Users\ADMIN\AppData\Local\Temp')
OUT = ROOT / 'Bao cao thuc hanh dich nguoc SandSort.docx'
PAGES = []

def crop(name, source, box):
    src = WORK / source if source.startswith('image') else TEMP / ('codex-clipboard-' + source + '.png')
    im = Image.open(src).convert('RGB')
    target = WORK / (name + '.png')
    im.crop(box).save(target)
    return str(target)

figs = {
 'package': crop('package','image2.png',(344,240,1131,834)),
 'asset': crop('asset','image1.png',(206,95,1248,940)),
 'export': crop('export','image4.png',(345,240,1130,510)),
 'dump': crop('dump','image9.png',(343,242,1130,805)),
 'import': crop('import','image8.png',(464,416,970,675)),
 'fingerprint': crop('fingerprint','image7.png',(359,143,1108,733)),
 'coinfield': crop('coinfield','1b3bfc54-e169-4ac0-958c-885532023f6d',(515,155,1170,700)),
 'getter': crop('getter','8fed4875-b0b1-4d8d-aaa5-34243197d5fc',(519,153,1175,304)),
 'setter': crop('setter','a654b0ca-a224-43f3-9c6b-2803974ae7a8',(519,154,1175,294)),
 'state_decl': crop('state_decl','41bdc8b5-3461-434f-b040-da3d5027ce97',(519,153,1178,300)),
 'dump_rva': crop('dump_rva','bd3743ee-ccf5-41ae-a712-9f57ab611ea5',(418,516,821,640)),
 'extension': crop('extension','63e8f8e5-250f-4819-a5e6-098d79a48580',(617,208,1320,755)),
 'scriptok': crop('scriptok','33b0a007-ece7-4f2b-a88e-ea70c0bbe436',(453,842,1130,901)),
 'native': crop('native','95159a74-0462-4426-bbe2-624871084082',(587,398,1440,747)),
 'start': crop('start','a621778a-9263-43c0-af36-4e28a77c7838',(1146,158,1919,650)),
 'key': crop('key','249a9e64-b87d-4010-aaba-67ca2daa2cef',(720,436,1559,612)),
 'awake': crop('awake','f23e6a62-3d87-4e41-becc-454dbb609759',(695,407,1508,567)),
 'state': crop('state','010cef5e-0bd7-44f2-a586-920cfcb57070',(1094,164,1918,586)),
 'state_save': crop('state_save','010cef5e-0bd7-44f2-a586-920cfcb57070',(1114,426,1832,481)),
 'next_wrapper': crop('next_wrapper','1bc355a1-a873-4556-b010-8eed68f3dd36',(1414,602,1910,748)),
 'next_call': crop('next_call','c18c8c8d-983c-4827-9cb1-fafa6e28695a',(1118,168,1918,558)),
 'call_asm': crop('call_asm','c18c8c8d-983c-4827-9cb1-fafa6e28695a',(594,524,1101,600)),
 'claim': crop('claim','38df2b30-d4b0-4445-b1e0-16a0b84adbae',(1225,171,1916,750)),
 'claimfield': crop('claimfield','2ccc4d0b-1c09-4f0c-ad7d-27f6688a9b63',(545,493,1000,643)),
 'reset': crop('reset','37c1c017-834c-4e5a-9ea9-cbe978919916',(1130,498,1917,621)),
}

def page(title, blocks):
    PAGES.append({'title': title, 'blocks': blocks})
def p(text): return ('p', text)
def h(text): return ('h', text)
def code(text): return ('code', text)
def fig(key, caption, width=16.8): return ('fig', key, caption, width)
def table(headers, rows, widths): return ('table', headers, rows, widths)

page('Báo cáo thực hành dịch ngược game Android SandSort', [
 p('Bài thực hành khảo sát cấu trúc APK Unity, xuất tài nguyên hình ảnh và phân tích hai luồng xử lý coin trong thư viện IL2CPP ARM64. Mục tiêu là xác định đúng lớp, phương thức, vị trí mã thực thi và đường đi của giá trị; không thực hiện sửa APK hoặc thay đổi số dư trong game.'),
 p('Kết quả phân tích tĩnh cho thấy PlayerResourceManager.Start nạp coin từ PlayerPrefs với khóa Coin_SandSort, dùng 100 khi chưa tồn tại khóa. Trong nhánh nhận thưởng của UIWinGame.NextLevel, phần thực thi MoveNext gọi StateCoin với state = true và coin = 100. StateCoin cập nhật field _coin rồi lưu số dư mới. Một tài nguyên Texture2D hình chai cũng đã được xuất bằng UABEA.'),
 h('1 Môi trường và dữ liệu thực hành'),
 p('Các thao tác được thực hiện trong Windows chạy trên VMware Workstation. Thư mục làm việc trong máy ảo là C:\\LabGame. Dữ liệu giải nén được đặt tại APK_giai_nen; tài nguyên xuất tại Asset_xuat; đầu ra dump nằm trong thư mục Il2CppDumper-win-v6.7.46 (1); project Ghidra có tên dichnguoc.'),
 table(['Thành phần','Thông tin và trạng thái'],[
 ['APK mẫu','SandSort theo đề thực hành; namespace game là FoodSort. Mã nhận diện thư viện được ghi ở mục 1.1.'],
 ['UABEA','UABEANext; đã mở dữ liệu Unity, xem Preview và xuất Texture2D. Workspace hiển thị Unity 2022.3.54f1.'],
 ['Il2CppDumper','Gói v6.7.46; đã tạo dump.cs, DummyDll và script.json.'],
 ['dnSpy','v6.1.8, 64 bit .NET; đã mở Assembly-CSharp.dll trong DummyDll và tra cứu các thành viên.'],
 ['Ghidra','v12.1.2; đã import ELF AARCH64 và chạy ghidra.py thành công sau khi bật Jython.'],
 ['Kiến trúc','arm64-v8a; AARCH64, little-endian, địa chỉ 64 bit.']], [3.1,13.7]),
 p('Các hình trong các mục sau chứng minh trạng thái mở và sử dụng công cụ. Không ghi thêm phiên bản VMware, cấu hình RAM/CPU hoặc phiên bản APK khi chưa có dữ liệu xác nhận.'),
])

page('1.1 Mã nhận diện thư viện được phân tích',[
 p('Thư viện native sử dụng xuyên suốt bài là C:\\LabGame\\APK_giai_nen\\lib\\arm64-v8a\\libil2cpp.so. Cửa sổ Import Results Summary của Ghidra ghi nhận định dạng ELF shared object, processor AARCH64 và mã nhận diện dưới đây. Mã băm này thuộc libil2cpp.so, không phải mã băm của toàn bộ APK.'),
 code('GNU Build ID\n02f85baf0795b27bc93c81d2846ad54135610db3\n\nMD5 của libil2cpp.so\nfbfcc2616006ee6ef7bc501ec0c6b086\n\nSHA256 của libil2cpp.so\ne07d562a09889db333a1f6e1d85bb74356cc29fa92f788a8154e16ac1b5cab3f'),
 p('Chưa xác minh versionName, versionCode hoặc tên file APK gốc. Vì vậy, bản thư viện native được nhận diện bằng Build ID và mã băm để người chấm đối chiếu, tránh nhầm với một lần dump hoặc một bản APK khác.'),
 fig('fingerprint','Thông tin kiến trúc và mã nhận diện của libil2cpp.so trong Ghidra',15.7),
])

page('2 Khảo sát APK và xuất tài nguyên Unity',[
 h('2.1 Cấu trúc APK sau khi giải nén'),
 p('Bản sao APK được giải nén vào C:\\LabGame\\APK_giai_nen. Các thư mục assets, lib, res và các file classes.dex, classes2.dex cùng AndroidManifest xuất hiện trong thư mục này. Thư mục assets chứa dữ liệu Unity; lib chứa thư viện native theo kiến trúc. Các file DEX và manifest thuộc phần ứng dụng Android, không thay thế mã IL2CPP của game.'),
 p('Trong UABEA, workspace data.unity3d chứa các mục sharedassets0.assets, sharedassets1.assets, sharedassets2.assets và các file dữ liệu liên quan. Dữ liệu Unity được khảo sát dưới assets/bin/Data. sharedassets0.assets là mục serialized asset bên trong workspace đang mở; không mặc định coi tên này là một file rời ngoài ổ đĩa nếu chưa kiểm tra cách đóng gói.'),
 fig('package','Cấu trúc thư mục APK_giai_nen sau khi giải nén',16.0),
])

page('2.2 Xem trước tài nguyên hình ảnh',[
 p('Trong workspace data.unity3d, chọn sharedassets0.assets và duyệt các đối tượng Texture2D. Asset được chọn có tên bottle 1, loại Texture2D. Khung Preview hiển thị một chai chứa vật liệu màu vàng cam; kích thước texture hiển thị là 281 × 552 pixel. Đây là tài nguyên hình ảnh, không phải phương thức xử lý coin.'),
 p('Theo hình dạng trong Preview, tài nguyên có thể phục vụ phần hiển thị chai chứa cát hoặc vật liệu màu trong nội dung game. Nhận xét này dựa trên hình ảnh tài nguyên; chưa truy vết prefab hoặc scene để khẳng định toàn bộ vị trí sử dụng.'),
 fig('asset','Asset bottle 1 thuộc loại Texture2D trong sharedassets0.assets và ảnh Preview',16.5),
])

page('2.3 Kết quả xuất asset',[
 p('Sau khi xem Preview, sử dụng chức năng Export của UABEA để xuất một bản sao tài nguyên vào C:\\LabGame\\Asset_xuat. File Explorer hiển thị một file ảnh có thumbnail hình chai tương ứng với Preview. Tên hiển thị của file là Bottle_Texture2D_sharedassets0.assets; phần mở rộng ảnh không hiện trong ảnh chụp nên không tự khẳng định định dạng PNG hoặc định dạng khác.'),
 fig('export','File tài nguyên đã xuất trong thư mục Asset_xuat'),
 p('Kết quả C1 gồm vị trí workspace gốc, mục sharedassets0.assets, tên bottle 1, loại Texture2D, ảnh Preview và bằng chứng file xuất. UABEA hỗ trợ nhận diện và trích xuất tài nguyên Unity; việc xác định nơi cộng coin phải được thực hiện bằng metadata và mã native ở các bước sau.'),
 h('3 Tạo dữ liệu tra cứu IL2CPP'),
 p('Il2CppDumper được sử dụng để đối chiếu mã máy trong libil2cpp.so với thông tin lớp và phương thức trong global-metadata.dat. Hai file đầu vào lấy từ cùng APK đã giải nén; kiến trúc ARM64 được giữ thống nhất khi import sang Ghidra.'),
 code('File mã native\nC:\\LabGame\\APK_giai_nen\\lib\\arm64-v8a\\libil2cpp.so\n\nFile metadata\nC:\\LabGame\\APK_giai_nen\\assets\\bin\\Data\\Managed\\Metadata\\global-metadata.dat'),
 p('Đường dẫn metadata là vị trí đầu vào theo cấu trúc IL2CPP được dùng trong bài. Các ảnh kết quả hiện thể hiện thư mục đầu ra và bản thư viện import; chưa có ảnh riêng của cả hai hộp chọn đầu vào Il2CppDumper.'),
])

page('3.1 Đầu ra và cách sử dụng dữ liệu dump',[
 p('Thư mục kết quả Il2CppDumper-win-v6.7.46 (1) có DummyDll, dump.cs, script.json, ghidra.py, il2cpp.h và stringliteral.json. Các file này phục vụ tra cứu, gắn tên và đối chiếu địa chỉ; không phải toàn bộ mã nguồn C# gốc của game.'),
 table(['Đầu ra','Vai trò trong bài'],[
 ['dump.cs','Tra tên lớp, field, chữ ký phương thức, RVA và File Offset.'],
 ['DummyDll','Mở Assembly-CSharp.dll trong dnSpy để xem cấu trúc namespace và thành viên.'],
 ['script.json','Cung cấp tên và địa chỉ để script gắn nhãn trong Ghidra.'],
 ['ghidra.py','Áp dụng dữ liệu script.json vào chương trình native đã import.'],
 ['il2cpp.h và stringliteral.json','Dữ liệu hỗ trợ về kiểu và chuỗi; chỉ sử dụng khi cần đối chiếu thêm.']], [4.0,12.8]),
 fig('dump','Các file đầu ra của Il2CppDumper trong thư mục làm việc',15.3),
])

page('4 Tìm ứng viên bằng dnSpy và dump.cs',[
 h('4.1 Xác định field chứa số coin'),
 p('Mở DummyDll\\Assembly-CSharp.dll trong dnSpy, duyệt namespace FoodSort và tìm từ khóa coin. Trong PlayerResourceManager tìm thấy private int _coin với FieldOffset 0x20. Kiểu int phù hợp với giá trị số lượng coin; vai trò của field được kiểm chứng tiếp bằng lệnh đọc và ghi trong mã native.'),
 p('Kết quả tìm kiếm cũng có _coinGain của UIWinGame, nhưng field này có kiểu Transform. Vì vậy, nó được xem là ứng viên tài nguyên giao diện, không dùng làm bằng chứng trực tiếp cho biến số dư. Tên field chỉ định hướng tìm kiếm; kết luận xử lý coin phải dựa trên đường đi của giá trị.'),
 fig('coinfield','Field _coin kiểu int trong FoodSort.PlayerResourceManager có offset 0x20',15.8),
])

page('4.2 Danh sách phương thức ứng viên và đối chiếu RVA',[
 p('Các phương thức get_Coin, set_Coin, Awake, Start và StateCoin được chọn để kiểm tra. Trong dnSpy, địa chỉ native được lấy từ thuộc tính Address, sau đó đối chiếu dòng RVA trong dump.cs. RVA ở chú thích của DLL giả và Token metadata không được dùng thay cho địa chỉ native.'),
 table(['Phương thức','RVA','File Offset'],[
 ['get_Coin','0x22A0708','0x229C708'],['set_Coin','0x22A0710','0x229C710'],['Awake','0x22A15D0','0x229D5D0'],['Start','0x22A15D4','0x229D5D4'],['StateCoin','0x22A16F8','0x229D6F8']], [6.5,5.15,5.15]),
 p('get_Coin trả về int, không có tham số; set_Coin nhận int value và trả về void. Thân hàm return 0 hoặc rỗng trong DummyDll là nội dung dựng phục vụ tra cứu, không chứng minh giá trị coin bằng 0 hoặc hàm không hoạt động.'),
 fig('getter','Khai báo get_Coin và thuộc tính Address trong dnSpy'),
 fig('dump_rva','RVA và File Offset của get_Coin và set_Coin trong dump.cs',14.5),
 p('StateCoin có chữ ký public void StateCoin(bool state, int coin, string placement = "unknown", string groupPlacement = "unknown"). Hướng cộng hoặc trừ của state chưa được suy đoán ở bước này; được xác nhận trong Ghidra ở mục 7.'),
])

page('5 Import và gắn tên phương thức trong Ghidra',[
 h('5.1 Import đúng thư viện ARM64'),
 p('Tạo project Non-Shared tên dichnguoc và import libil2cpp.so đã dùng trong lần dump. Hộp import nhận dạng Executable and Linking Format (ELF), language AARCH64:LE:64:v8A:default. Sau khi mở CodeBrowser, thực hiện phân tích tự động để nhận diện các vùng mã, hàm và tham chiếu.'),
 fig('import','Import libil2cpp.so với định dạng ELF và ngôn ngữ AARCH64',15.7),
 h('5.2 Quy ước địa chỉ'),
 p('Ghidra sử dụng Image Base 0x100000 trong lần import này. Ví dụ, get_Coin có RVA 0x22A0708 và được gắn nhãn tại 0x23A0708. File Offset là vị trí trong file ELF, còn địa chỉ Ghidra là vị trí trong không gian địa chỉ của chương trình đã import; hai giá trị không được tráo đổi.'),
 code('Địa chỉ Ghidra = Image Base + RVA\n0x100000 + 0x22A0708 = 0x23A0708'),
 p('Listing cung cấp lệnh ARM64 và địa chỉ; Decompiler dựng mã giả để giải thích luồng xử lý. Mã giả được đối chiếu với khai báo dnSpy và các offset, không coi là mã C++ hoặc C# gốc.'),
])

page('5.3 Khắc phục môi trường chạy script và hoàn tất gắn tên',[
 p('Thêm thư mục Il2CppDumper-win-v6.7.46 (1) vào Script Directories. Chọn cả thư mục, không thêm riêng file ghidra.py làm bundle. Lần chạy đầu báo Ghidra was not started with PyGhidra. Python is not available, cho thấy script chưa được thực thi bằng môi trường tương thích.'),
 p('Trong cửa sổ quản lý project, mở File > Install Extensions, bật Jython rồi khởi động lại Ghidra. Do script Il2CppDumper sử dụng kiểu Python 2/Jython, thêm chỉ thị runtime vào đầu ghidra.py để Ghidra chọn đúng môi trường:'),
 code('# -*- coding: utf-8 -*-\n# @runtime Jython\nimport json'),
 fig('extension','Extension Jython phiên bản 12.1.2 được bật trong Ghidra',13.8),
 p('Chạy lại ghidra.py và chọn script.json của cùng lần dump. Console xuất hiện Script finished! và ghidra.py> Finished!. Sau đó lưu project; tên phương thức FoodSort xuất hiện trong Symbol Tree.'),
 fig('scriptok','Thông báo hoàn tất của lần chạy script thành công'),
])

page('6 Xác định nơi đọc và ghi field coin trong mã native',[
 p('get_Coin tại 0x23A0708 chỉ gồm lệnh ldr w0, [x0, #0x20] rồi ret. Lệnh đọc giá trị 32 bit từ offset 0x20 của đối tượng và trả về; khớp private int _coin trong dnSpy. set_Coin tại 0x23A0710 thực hiện str w1, [x0, #0x20] rồi ret, ghi giá trị truyền vào cùng field.'),
 fig('native','Listing của get_Coin và set_Coin xác nhận đọc và ghi offset 0x20'),
 p('Kết quả này xác định vị trí field trong đối tượng PlayerResourceManager. Offset 0x20 không phải số coin, RVA hoặc địa chỉ cố định của một đối tượng khi game chạy. Kiểu undefined4 trong Decompiler chỉ phản ánh việc Ghidra chưa khôi phục tên kiểu cho dữ liệu 4 byte; dnSpy cho biết field là int.'),
 p('Hai hàm accessor không chứa phép cộng hoặc lời gọi PlayerPrefs trong đoạn đã xem. Vì vậy, chúng là điểm đối chiếu field, không được kết luận là nơi tính thưởng hoặc nạp dữ liệu ban đầu. Việc nạp coin được tìm trong Start; phép thay đổi số dư được tìm trong StateCoin.'),
 h('6.1 Loại trừ Awake khỏi luồng nạp coin'),
 p('Awake tại 0x23A15D0 chỉ có lệnh ret. Đây là bằng chứng native để loại phương thức này khỏi ứng viên trực tiếp đọc hoặc khởi tạo coin trong bản thư viện đang phân tích.'),
 fig('awake','Awake của PlayerResourceManager chỉ trả về, không nạp coin'),
])

page('6.2 Logic coin khi vào game',[
 p('Hàm xác định là FoodSort.PlayerResourceManager.Start, RVA 0x22A15D4, địa chỉ Ghidra 0x23A15D4. Ở dòng 18, hàm gọi UnityEngine.PlayerPrefs.GetInt với khóa _StringLiteral_3130 và giá trị mặc định 100. Dòng 19 ghi kết quả vào param_1 + 0x20, chính là field _coin đã đối chiếu.'),
 fig('start','Start đọc coin từ PlayerPrefs rồi gán vào field _coin',16.5),
 p('Luồng giá trị là dữ liệu đã lưu → GetInt → _coin. Nếu khóa tồn tại, số coin được lấy từ lần lưu trước; nếu chưa có khóa, giá trị mặc định là 100. Do đó, không kết luận mọi lần vào game đều nhận lại 100 coin. Các lời gọi GetInt khác ghi vào offset 0x24, 0x28 và 0x2C thuộc các field booster, không gộp chúng vào biến coin.'),
])

page('6.3 Đối chiếu khóa lưu và kết luận luồng khởi tạo',[
 p('Đi theo _StringLiteral_3130 trong Listing cho thấy chú thích Coin_SandSort ở nhãn StringLiteral_3130. Đây là khóa được truyền vào GetInt trong Start và SetInt trong StateCoin. Việc hai phương thức dùng cùng khóa nối kết bước nạp và bước lưu số dư.'),
 fig('key','Nhãn StringLiteral_3130 mang nội dung Coin_SandSort'),
 p('Mã diễn giải tương đương cho phần nạp coin được viết như sau. Đây là cách diễn đạt hành vi đã quan sát, không phải mã nguồn được khôi phục nguyên bản:'),
 code('_coin = PlayerPrefs.GetInt("Coin_SandSort", 100);'),
 p('Kết luận mục 6: Start thực sự nạp hoặc khởi tạo coin; nguồn dữ liệu là PlayerPrefs phía client; giá trị mặc định là 100; đích ghi là field int _coin, offset 0x20. get_Coin/set_Coin xác nhận nơi đọc và ghi field. Awake được loại trừ vì mã native chỉ trả về.'),
 h('7 Logic cập nhật coin sau khi nhận thưởng'),
 p('Phần tiếp theo phân biệt nơi thực hiện cộng/trừ số dư với nơi quyết định phần thưởng. PlayerResourceManager.StateCoin nhận lượng coin cần thay đổi; nhánh trong UIWinGame.NextLevel quyết định truyền giá trị 100 và chọn hướng cộng. Phân tích chỉ kết luận nhánh đã quan sát, không áp dụng mặc định cho tất cả nguồn thưởng trong game.'),
])

page('7.1 Hàm StateCoin cộng trừ và lưu số dư',[
 p('FoodSort.PlayerResourceManager.StateCoin có RVA 0x22A16F8, địa chỉ Ghidra 0x23A16F8. Đối chiếu khai báo dnSpy: param_1 là đối tượng; param_2 tương ứng bool state; param_3 tương ứng int coin. Trong ảnh, tham số state và coin được lưu vào đối tượng phụ trước khi sử dụng, nhưng field số dư vẫn nằm ở param_1 + 0x20.'),
 p('Dòng 39 khởi tạo số thay đổi bằng -coin. Dòng 40–41 chuyển thành +coin nếu state khác 0. Dòng 43 cộng số thay đổi vào coin hiện tại; dòng 44 ghi lại _coin. Dòng 45 gọi PlayerPrefs.SetInt với khóa _StringLiteral_3130; dòng 46 gọi PlayerPrefs.Save.'),
 fig('state','Nhánh cộng trừ, cập nhật field và lưu coin trong StateCoin'),
 code('int change = state ? coin : -coin;\n_coin = _coin + change;\nPlayerPrefs.SetInt("Coin_SandSort", _coin);\nPlayerPrefs.Save();'),
 p('Với coin là số dương, state = true làm tăng số dư và state = false làm giảm số dư. Phần mã tiếp theo có RecordLevelCoinChange và TrackWhenMetadataReady, nhưng chưa phân tích đầy đủ cơ chế thống kê; không dùng tên những hàm này để suy ra máy chủ xác thực số dư.'),
])

page('7.2 Nhánh thưởng 100 coin trong luồng NextLevel',[
 p('UIWinGame.NextLevel tại 0x23B6044, RVA 0x22B6044, chuẩn bị state machine <NextLevel>d__26 và gọi AsyncVoidMethodBuilder.Start. Đoạn nghiệp vụ được quan sát trong <NextLevel>d__26.MoveNext. Trong nhánh này, TryStartClaim phải trả về true trước khi lấy đối tượng PlayerResourceManager và gọi StateCoin.'),
 p('Dòng 42–43 của ảnh dưới truyền state = 1 và coin = 100. Listing xác nhận mov w1, #0x1 và mov w2, #0x64; 0x64 tương đương 100. Địa chỉ lệnh gọi là 0x23B7804, không được nhầm với địa chỉ bắt đầu NextLevel hoặc MoveNext.'),
 fig('next_call','MoveNext kiểm tra TryStartClaim rồi gọi StateCoin với lượng thưởng 100'),
 fig('call_asm','Lệnh ARM64 thiết lập tham số và gọi StateCoin tại 0x23B7804'),
 p('Hai tham số chuỗi ở lời gọi mang nhãn _StringLiteral_13854 và _StringLiteral_14787; chưa xác định nội dung nên không tự đặt tên nguồn thưởng. Giá trị 100 ở đây là lượng cộng, khác với 100 mặc định khi GetInt chưa tìm thấy khóa trong Start.'),
])

page('7.3 Điều kiện bắt đầu nhận thưởng',[
 p('FoodSort.UIWinGame.TryStartClaim có địa chỉ Ghidra 0x23B632C, RVA 0x22B632C. Hàm đọc cờ ở offset 0xA8 vào cVar1. Nếu giá trị ban đầu bằng 0, nó đặt cờ thành 1 và thực hiện các thao tác chuẩn bị. Cuối hàm trả về cVar1 == 0, dựa trên giá trị cờ ban đầu.'),
 fig('claim','TryStartClaim kiểm tra cờ nhận thưởng và trả về kết quả điều kiện',14.0),
 p('Như vậy, cờ ban đầu false cho phép nhánh cộng coin đi tiếp; cờ đang true làm nhánh này bị bỏ qua. Kết luận áp dụng trong trường hợp hàm hoàn tất bình thường; các nhánh gặp đối tượng null có thể đi vào xử lý lỗi. Hai lần gọi liên tiếp khi cờ chưa được reset không đều nhận thưởng.'),
])

page('7.4 Đối chiếu tên cờ và nơi đặt lại',[
 p('Trong dnSpy, field private bool _isClaimingReward có FieldOffset 0xA8, khớp vị trí TryStartClaim đã đọc. Field _isRewardAdCallbackHandled ở offset 0xA9 là một cờ khác; không gộp hai cờ thành một điều kiện.'),
 fig('claimfield','Field _isClaimingReward ở offset 0xA8 trong UIWinGame',14.2),
 p('UIWinGame.OnEnable có RVA 0x22B5B40, địa chỉ Ghidra 0x23B5B40. Trong nhánh UIManager tồn tại, dòng 71 ghi 0 vào offset 0xA8 và dòng 72 gọi BindClaimActions. Điều này cho thấy cờ nhận thưởng được đặt lại khi đi qua nhánh bật giao diện này. Không suy ra cờ chỉ được reset một lần cho mỗi level hoặc đây là toàn bộ cơ chế chống nhận lặp.'),
 fig('reset','OnEnable đặt lại cờ _isClaimingReward và gọi BindClaimActions',16.8),
])

page('7.5 Tổng hợp luồng và giới hạn chứng minh',[
 p('Chuỗi xử lý đã xác định là UIWinGame.NextLevel → phần thực thi <NextLevel>d__26.MoveNext → TryStartClaim → PlayerResourceManager.StateCoin(true, 100, ...) → cập nhật _coin → PlayerPrefs.SetInt("Coin_SandSort", số dư mới) → PlayerPrefs.Save. OnEnable cung cấp bằng chứng đặt lại cờ _isClaimingReward trong nhánh đã quan sát.'),
 p('NextLevel không chứa trực tiếp phép cộng vì đây là hàm bất đồng bộ. Việc truy vết sang MoveNext giúp tìm đúng mã nghiệp vụ thay vì kết luận dựa vào phần tạo AsyncVoidMethodBuilder.'),
 table(['Vị trí','Địa chỉ Ghidra','Ý nghĩa'],[
 ['NextLevel','0x23B6044','Khởi chạy luồng bất đồng bộ'],
 ['TryStartClaim','0x23B632C','Kiểm tra cờ bắt đầu nhận thưởng'],
 ['Lời gọi StateCoin','0x23B7804','Truyền state = true và coin = 100'],
 ['StateCoin','0x23A16F8','Cộng/trừ và lưu số dư'],
 ['OnEnable','0x23B5B40','Đặt lại cờ trong nhánh đã xem']], [5.0,4.2,7.6]),
 p('Chưa xác định đầy đủ sự kiện báo thắng hoặc nút cụ thể nào gọi NextLevel. XREF của NextLevel trong ảnh hiện cho thấy tham chiếu dữ liệu, chưa chứng minh lời gọi trực tiếp từ sự kiện hoàn thành level. Do đó, kết luận là tìm được nhánh thưởng trong lớp giao diện thắng UIWinGame, không khẳng định mọi lần thắng đều tự động cộng ngay 100 coin.'),
 h('Các ứng viên đã loại trừ hoặc thu hẹp vai trò'),
 table(['Ứng viên','Kết quả đối chiếu và lý do'],[
 ['PlayerResourceManager.Awake','Chỉ có ret tại 0x23A15D0; loại khỏi nơi trực tiếp nạp hoặc cộng coin.'],
 ['get_Coin và set_Coin','Chỉ đọc/ghi _coin ở +0x20; không có phép cộng hay lưu PlayerPrefs trong đoạn đã xem.'],
 ['UIWinGame._coinGain','Có kiểu Transform; là thành phần giao diện, không phải field int chứa số dư.'],
 ['PlayerResourceManager.Start','Đọc số dư đã lưu hoặc dùng mặc định 100; không phải bằng chứng cộng thưởng sau thắng.'],
 ['UIWinGame.NextLevel','Giữ lại làm đầu mối nhận thưởng; phần thân chỉ khởi chạy state machine, cần đọc MoveNext.']], [5.2,11.6]),
])

page('8 Nhận xét an toàn và giới hạn phân tích tĩnh',[
 h('8.1 Kịch bản can thiệp tài nguyên phía client'),
 p('Một kịch bản có thể xảy ra là thay đổi dữ liệu lưu cục bộ của khóa Coin_SandSort trước khi game nạp lại. Nếu thay đổi có hiệu lực và không bị cơ chế khác kiểm tra hoặc ghi đè, Start sẽ đọc số đã thay đổi rồi đưa vào _coin. Kịch bản thứ hai là can thiệp lượng thưởng truyền vào StateCoin hoặc thao tác cập nhật số dư, từ đó làm tăng tài nguyên ngoài mức mà nhánh thưởng thông thường quyết định.'),
 p('Nhận định có cơ sở từ ba bằng chứng: nguồn đọc là PlayerPrefs; phép cộng diễn ra trong libil2cpp.so phía client; số dư mới được ghi lại cục bộ. Trong bài chỉ xác định khả năng và vị trí liên quan, chưa sửa APK, sửa dữ liệu lưu, chạy bản can thiệp hoặc đo tác động trong game.'),
 h('8.2 Tác động và giới hạn của kết luận'),
 p('Tác động dự kiến là số dư và phần hiển thị tài nguyên phía client có thể sai lệch, ảnh hưởng các thao tác tiêu coin hoặc cân bằng game nếu chúng tin vào số dư này. Tuy nhiên, các ảnh không chứng minh rằng server chấp nhận số dư sai, không có xác thực từ xa, hoặc người chơi có thể dùng số coin đó cho mọi chức năng.'),
 p('Phân tích tĩnh chưa xác minh trạng thái lúc chạy, dữ liệu PlayerPrefs thực tế, kết quả hiển thị sau thắng, nội dung hai chuỗi mô tả nguồn thưởng hoặc toàn bộ các nhánh thưởng khác. Các hàm thống kê xuất hiện trong StateCoin cũng chưa được khảo sát đầy đủ. Vì vậy, không kết luận đã khai thác thành công hoặc game hoàn toàn không có cơ chế bảo vệ.'),
 h('8.3 Cách giảm phụ thuộc vào dữ liệu do client quyết định'),
 p('Với game có tài khoản và máy chủ, nên để máy chủ quản lý số dư có giá trị, kiểm tra điều kiện hoàn thành level và xác định lượng thưởng hợp lệ. Mỗi giao dịch thưởng cần mã nhận diện để tránh xử lý lặp; client gửi yêu cầu và nhận số dư được xác nhận thay vì tự quyết định phần thưởng. Dữ liệu local có thể dùng làm bộ nhớ đệm để hiển thị, sau đó đối chiếu với trạng thái được máy chủ xác nhận.'),
 p('Trong chế độ offline, kiểm tra tính toàn vẹn và đối chiếu dữ liệu cục bộ có thể tăng khó khăn cho việc sửa trực tiếp, nhưng vẫn không tạo được mức tin cậy tương đương một nguồn trạng thái độc lập với client. Cờ _isClaimingReward hữu ích để hạn chế thao tác nhận lặp trong luồng giao diện, không phải bằng chứng xác thực số dư từ máy chủ.'),
])

page('9 Kết luận và kết quả cần nộp',[
 p('Bài thực hành đã nhận diện dữ liệu Unity trong APK, chọn Texture2D bottle 1 trong sharedassets0.assets, xem Preview và xuất tài nguyên. Il2CppDumper tạo dữ liệu tra cứu, dnSpy cung cấp cấu trúc lớp và offset, còn Ghidra xác nhận hành vi native. Việc bật Jython và chỉ định runtime giúp ghidra.py gắn tên thành công trên Ghidra 12.1.2.'),
 table(['Nội dung','Kết quả xác định'],[
 ['Nạp coin khi vào game','PlayerResourceManager.Start, RVA 0x22A15D4, địa chỉ 0x23A15D4; đọc Coin_SandSort, mặc định 100, ghi _coin +0x20.'],
 ['Cập nhật số dư','StateCoin, RVA 0x22A16F8, địa chỉ 0x23A16F8; cộng/trừ theo state, ghi field rồi lưu PlayerPrefs.'],
 ['Nhánh thưởng đã xem','NextLevel.MoveNext gọi StateCoin(true, 100, ...) tại 0x23B7804 khi TryStartClaim cho phép.'],
 ['Điều kiện nhận thưởng','Cờ _isClaimingReward tại +0xA8; được kiểm tra và đặt lại trong các nhánh đã phân tích.'],
 ['Phần chưa xác định','Đường gọi từ sự kiện hoàn thành level/nút cụ thể đến NextLevel; kiểm chứng động và vai trò server.']], [4.2,12.6]),
 h('9.1 Thư mục kết quả'),
 p('Kết quả kèm báo cáo gồm file asset đã xuất trong Asset_xuat; dump.cs, DummyDll, script.json và ghidra.py đã thêm runtime Jython; project Ghidra đã lưu và ảnh minh chứng. APK gốc và thư mục APK_giai_nen được giữ để đối chiếu theo hướng dẫn GVTH. Khi nộp, bổ sung tên file APK hoặc thông tin phiên bản nếu được giảng viên yêu cầu, đồng thời giữ mã nhận diện thư viện ở mục 1.1.'),
 h('9.2 Tài liệu sử dụng'),
 p('1. Tài liệu thực hành của giảng viên: Thực hành dịch ngược, bài Kỹ thuật dịch ngược game Android, các phần B, C và D.'),
 p('2. Il2CppDumper: mã script ghidra.py và chức năng các đầu ra tra cứu. https://github.com/Perfare/Il2CppDumper'),
 p('3. Ghidra: hướng dẫn cài extension và lịch sử thay đổi cơ chế chọn runtime Python. https://github.com/NationalSecurityAgency/ghidra'),
])

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(1.8)
sec.left_margin=sec.right_margin=Cm(2.1)
normal=doc.styles['Normal']
normal.font.name='Times New Roman'; normal.font.size=Pt(11.5)
normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.widow_control=True
for s in ['Title','Heading 1','Heading 2','Caption']:
    st=doc.styles[s]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.space_before=Pt(3)
doc.styles['Title'].font.size=Pt(20)
doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Heading 2'].font.size=Pt(12.5)
doc.styles['Caption'].font.size=Pt(10)
doc.styles['Caption'].font.italic=True
doc.styles['Caption'].paragraph_format.keep_with_next=False
doc.styles['Caption'].paragraph_format.space_after=Pt(8)
doc.core_properties.title='Báo cáo thực hành dịch ngược game Android SandSort'
doc.core_properties.subject='Phân tích asset Unity và hai luồng coin trong IL2CPP ARM64'
doc.core_properties.author=''
doc.core_properties.keywords='SandSort, IL2CPP, coin, UABEA, Ghidra'

def word_table(block):
    _,headers,rows,widths=block
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    for i,header in enumerate(headers):t.rows[0].cells[i].text=header
    for row in rows:
        cells=t.add_row().cells
        for i,text in enumerate(row):cells[i].text=text
    for rowi,row in enumerate(t.rows):
        for i,c in enumerate(row.cells):
            c.width=Cm(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for e in ['top','left','bottom','right']:
                el=OxmlElement('w:'+e);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            if rowi==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF4');tcpr.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3);para.paragraph_format.line_spacing=1.04
                for run in para.runs:run.font.size=Pt(10.5);run.bold=(rowi==0)
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

fig_num=0
for pi,pageobj in enumerate(PAGES):
    if pi:doc.add_page_break()
    doc.add_paragraph(pageobj['title'],style='Title' if pi==0 else 'Heading 1')
    for block in pageobj['blocks']:
        kind=block[0]
        if kind=='p':
            para=doc.add_paragraph(block[1]);para.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind=='h':doc.add_paragraph(block[1],style='Heading 2')
        elif kind=='code':
            para=doc.add_paragraph();para.paragraph_format.keep_together=True
            para.paragraph_format.space_before=Pt(3);para.paragraph_format.space_after=Pt(8)
            for j,line in enumerate(block[1].split('\n')):
                if j:para.add_run().add_break()
                run=para.add_run(line);run.font.name='Consolas';run.font.size=Pt(9)
        elif kind=='table':word_table(block)
        elif kind=='fig':
            fig_num+=1
            para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.keep_with_next=True;para.paragraph_format.space_after=Pt(3)
            run=para.add_run();run.add_picture(figs[block[1]],width=Cm(block[3]))
            inline=run._r.find('.//'+qn('wp:docPr'))
            if inline is not None:inline.set('descr',block[2])
            caption=doc.add_paragraph(f'Hình {fig_num}. {block[2]}',style='Caption');caption.alignment=WD_ALIGN_PARAGRAPH.CENTER

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
doc.save(OUT)
(WORK/'report_manifest.json').write_text(json.dumps({'pages':PAGES,'figures':figs,'output':str(OUT)},ensure_ascii=False,indent=2),encoding='utf-8')
print(f'OUTPUT {OUT}\nPLANNED_PAGES {len(PAGES)}\nFIGURES {fig_num}')
