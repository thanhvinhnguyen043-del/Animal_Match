from pathlib import Path
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import re, json, hashlib

root=Path(__file__).resolve().parent
source=root.parent/'Báo cáo dịch ngược SandSort cập nhật luồng thắng.docx'
target=root.parent/'Báo cáo dịch ngược SandSort rà soát ảnh và nội dung.docx'
work=root/'compact_review'
source_hash=hashlib.sha256(source.read_bytes()).hexdigest()
d=Document(source)
removed={18,27,31,39}
mapping={old:old-sum(n<old for n in removed) for old in range(1,47) if old not in removed}

def replace(p,text):
    if p.runs:
        p.runs[0].text=text
        for r in p.runs[1:]:r.text=''
    else:p.add_run(text)

def remove(el):
    el.getparent().remove(el)

deleted=[]
for p in list(d.paragraphs):
    match=re.match(r'^Hình (\d+)\.',p.text)
    if match and int(match.group(1)) in removed:
        old=int(match.group(1));prev=p._p.getprevious()
        # Captions in the source follow their image directly.
        assert prev is not None and prev.tag==qn('w:p') and prev.xpath('.//w:drawing')
        deleted.append({'old_number':old,'caption':p.text})
        remove(prev);remove(p._p)

for p in list(d.paragraphs):
    t=p.text
    if t=='Start đọc số coin từ khóa Coin_SandSort và ghi vào field _coin. Khi chưa có dữ liệu lưu, hàm dùng giá trị mặc định 100.':
        remove(p._p)
    elif t.startswith('Hình 27 hiển thị điểm bắt đầu StateCoin'):
        replace(p,'Hình 28 cho thấy phép tính cộng trừ và bước lưu số dư trong StateCoin.')
    elif t.startswith('Dòng 42 - 43 trong Hình 30'):
        replace(p,'Dòng 42–43 trong Hình 30 truyền state = 1 và coin = 100 vào StateCoin. Giá trị này được đối chiếu trong Listing bằng các lệnh mov w1, #0x1 và mov w2, #0x64; 0x64 tương đương 100. Lệnh gọi StateCoin nằm tại 0x23B7804, không phải địa chỉ bắt đầu NextLevel hoặc MoveNext.')
    elif t.startswith('OnPointerUp gọi callback qua delegate nên'):
        replace(p,'OnPointerUp thực thi callback qua delegate nên tại dòng gọi không xuất hiện trực tiếp tên NextLevel.')
    elif t.startswith('Luồng đã được nối từ kiểm tra thắng, mở UIWinGame'):
        replace(p,'Lượng thưởng 100 thuộc nhánh nhận thưởng qua nút tiếp tục, không được cộng tự động ngay khi CheckWin trả về true.')
    elif t=='Các ứng viên đã loại trừ hoặc thu hẹp vai trò':
        following=p._p.getnext()
        assert following.tag==qn('w:tbl')
        remove(following);remove(p._p)
    elif t.startswith('Start đọc số dư từ PlayerPrefs, còn StateCoin cập nhật'):
        replace(p,'Chưa có thử nghiệm can thiệp để kiểm tra tác động thực tế của hai kịch bản này.')
    elif t.startswith('Dữ liệu PlayerPrefs khi game chạy và số coin hiển thị'):
        replace(p,'Dữ liệu PlayerPrefs khi game chạy và số coin hiển thị sau khi nhận thưởng chưa được kiểm tra. Các nhánh của CheckWin và một số field chưa được đối chiếu đầy đủ; nhánh X2 mới được xác định ở bước gắn callback. Nội dung hai chuỗi mô tả nguồn thưởng và các hàm thống kê trong StateCoin cũng chưa được phân tích hết. Vì vậy, chưa đủ căn cứ để kết luận mọi lần thắng đều tự động cộng ngay 100 coin hoặc có thể can thiệp thành công vào số dư.')
    elif t.startswith('Với SandSort, đã xác định được cách nạp số coin'):
        replace(p,'Với SandSort, đã xác định được cách nạp số coin khi vào game, giá trị mặc định 100 khi chưa có dữ liệu lưu và cách cộng trừ, lưu số dư. Luồng kiểm tra thắng cũng đã được nối đến thao tác mở UIWinGame và nhận 100 coin qua nút tiếp tục. Kết quả cho thấy cần đọc cả MoveNext của phương thức bất đồng bộ và đối chiếu tên field để hiểu đúng mã giả. Kết quả dựa trên phân tích tĩnh và chưa được đối chiếu bằng việc chạy game.')

# Update every caption and prose reference with the same old-to-new mapping.
for el in d.element.xpath('.//w:p'):
    p=Paragraph(el,d)
    t=p.text
    if re.search(r'Hình\s+\d+',t):
        nums=[int(m.group(1)) for m in re.finditer(r'Hình\s+(\d+)',t)]
        assert not any(n in removed for n in nums),(t,nums)
        replace(p,re.sub(r'Hình\s+(\d+)',lambda m:f'Hình {mapping[int(m.group(1))]}',t))
    is_image=bool(el.xpath('.//w:drawing | .//w:pict'))
    is_caption=bool(re.match(r'^Hình\s+\d+\.',p.text))
    if is_image or is_caption:
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.left_indent=Pt(0)
        p.paragraph_format.right_indent=Pt(0)
    if is_caption:
        p.paragraph_format.keep_together=True
        p.paragraph_format.keep_with_next=False
        for r in p.runs:
            r.bold=True;r.italic=True;r.font.color.rgb=RGBColor(0,0,0)
    if is_image:p.paragraph_format.keep_with_next=True
    if p.style.name.startswith('Heading'):
        p.paragraph_format.page_break_before=False
for name in ('Heading 1','Heading 2','Heading 3'):
    d.styles[name].paragraph_format.page_break_before=False

# Remove orphaned image relationships so deleted figures are not embedded in the output.
used={el.get(qn('r:embed')) for el in d.element.xpath('.//a:blip')}
for rid,rel in list(d.part.rels.items()):
    if rel.reltype.endswith('/image') and rid not in used:d.part.drop_rel(rid)
d.save(target)
assert hashlib.sha256(source.read_bytes()).hexdigest()==source_hash
check=Document(target)
caps=[p for p in check.paragraphs if re.match(r'^Hình\s+\d+\.',p.text)]
assert len(caps)==42
assert [int(re.match(r'^Hình (\d+)',p.text).group(1)) for p in caps]==list(range(1,43))
assert len(check.inline_shapes)==43 # 42 report figures plus the cover logo
assert len(check.tables)==5
for old in range(1,9):
    assert mapping[old]==old
    oldp=next(p for p in Document(source).paragraphs if p.text.startswith(f'Hình {old}.'))
    newp=next(p for p in check.paragraphs if p.text.startswith(f'Hình {old}.'))
    assert oldp.text==newp.text
(work/'edit_manifest.json').write_text(json.dumps({'source':str(source),'source_sha256':source_hash,'output':str(target),'removed':deleted,'mapping':mapping,'figures':42},ensure_ascii=False,indent=2),encoding='utf8')
print('OUTPUT',target)
print('42 figures; initial tool-launch figures preserved; numbering and prose references updated.')
