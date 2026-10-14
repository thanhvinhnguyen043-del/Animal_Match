from pathlib import Path
import hashlib, json, re
from PIL import Image
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT=Path(__file__).resolve().parent
WORK=ROOT/'bento_report_expanded';WORK.mkdir(exist_ok=True)
REF=ROOT.parent/'Báo cáo thực hành dịch ngược BentoLoop hoàn chỉnh.docx'
OUT=ROOT.parent/'Báo cáo BentoLoop làm rõ phần 5 6 7.docx'
TEMP=Path(r'C:\Users\ADMIN\AppData\Local\Temp')
ref_hash=hashlib.sha256(REF.read_bytes()).hexdigest()
d=Document(REF)
original=(ROOT/'build_bento_report.py').read_text(encoding='utf-8-sig')
helpers=original[original.index('def para(text):'):original.index('d.add_page_break()')]
exec(compile(helpers,str(ROOT/'build_bento_report.py'),'exec'))
figures=[None]*7

body=d.element.body
start=next(p._p for p in d.paragraphs if p.text=='5 Phân tích điều kiện chấp nhận license')
anchor=next(p._p for p in d.paragraphs if p.text=='8 Build APK và kiểm thử bản đã sửa')
remove=[]
node=start
while node is not anchor:
    remove.append(node);node=node.getnext()
for node in remove:body.remove(node)
baseline=set(body)

def oldfig(n,caption,width=16):
    path=ROOT/'bento_report'/f'figure-{n:02d}.png'
    im=Image.open(path)
    box=(0,0,im.width,650) if n==16 else (0,0,im.width,im.height)
    return fig(path,box,caption,width)

def screenshot(uid,box,caption,width=16):
    source=TEMP/f'codex-clipboard-{uid}.png'
    saved=WORK/'evidence'/source.name;saved.parent.mkdir(exist_ok=True)
    saved.write_bytes(source.read_bytes())
    return fig(saved,box,caption,width)

heading('5 Phân tích điều kiện chấp nhận license')
heading('5.1 Xác định chuỗi được đem so sánh trong VerifyLicense',2)
para('Mở VerifyLicense tại 0x1F7689C bằng Go To và đọc phần xử lý đầu vào trong Decompiler. Đoạn cần theo dõi là các lời gọi IsNullOrWhiteSpace, Trim và Equals trong ảnh dưới.')
screenshot('94d8ddd2-1581-4280-bd08-3825a4ab5540',(1068,499,1914,682),'Các dòng xử lý đầu vào và lời gọi String.Equals với StringLiteral_2770',16.8)
para('Trong Hình 8, param_1 là chuỗi được truyền vào hàm. Dòng 19 gọi System.String$$IsNullOrWhiteSpace để kiểm tra chuỗi rỗng hoặc chỉ có khoảng trắng. Điều kiện ở dòng 20 chỉ đưa luồng vào phần kiểm tra tiếp khi kết quả bằng 0. Dòng 25 gọi System.String$$Trim và lưu chuỗi đã bỏ khoảng trắng ở hai đầu vào uVar3.')
para('Dòng 26 là vị trí cần chú ý: System.String$$Equals(uVar3, StringLiteral_2770, 4, 0). Chuỗi vừa xử lý được so sánh với StringLiteral_2770, còn kết quả so sánh được lưu vào uVar2. StringLiteral_2770 là nhãn dữ liệu chuỗi trong Ghidra, không phải một hàm. Ở bước này, đầu mối thu được là nhãn của chuỗi cần đọc để tìm key.')
heading('5.2 Theo nhánh điều kiện đến giá trị trả về',2)
para('Sau phép Equals, đọc tiếp các nhánh if và phần cuối VerifyLicense để xác định chuỗi khớp có thực sự dẫn đến kết quả thành công hay không.')
screenshot('b90788c8-59bd-4dba-a17e-6a7a2a451e6c',(1068,568,1916,939),'Phần cuối VerifyLicense gán kết quả thành công bằng 1 và thất bại bằng 0',16.8)
para('Khi phép so sánh trực tiếp không khớp, VerifyLicense gọi CheckTransformedLicense và có thể gọi tiếp SecurityAuthService.DispatchVerification. Nếu các nhánh bổ sung đều không chấp nhận đầu vào, luồng chuyển tới LAB_01f76914. Trong Hình 9, nhãn này gán uVar3 = 0 rồi đi đến return uVar3, tương ứng kết quả false.')
para('Nhánh so sánh trực tiếp khớp đi đến phần ghi log và phép gán uVar3 = 1 ở dòng 51. Dòng 57 trả về giá trị này, tương ứng true. Vì vậy, việc lần đến StringLiteral_2770 dựa trên cả lời gọi so sánh ở Hình 8 và đường trả về thành công ở Hình 9, chứ không chỉ dựa vào tên của nhãn.')
heading('5.3 Rút gọn luồng kiểm tra để chọn bước tiếp theo',2)
para('Các điều kiện quan sát được có thể viết lại như sau. Phần khởi tạo dữ liệu và ghi log được lược đi để thấy rõ thứ tự các nhánh chấp nhận đầu vào.')
code('if (IsNullOrWhiteSpace(inputKey)) return false;\nkey = Trim(inputKey);\nif (Equals(key, StringLiteral_2770, 4)) return true;\nif (CheckTransformedLicense(key)) return true;\nreturn SecurityAuthService.DispatchVerification(key);')
para('Luồng trên cho thấy cần đọc nội dung StringLiteral_2770 trước. Nhánh CheckTransformedLicense còn kiểm tra từng ký tự với mảng byte bằng XOR 0x53 và ((i × 13) AND 0xFF). Tuy nhiên, nhánh Equals đã chỉ ra một chuỗi có thể được chấp nhận trực tiếp, nên bước tiếp theo tập trung vào nhãn 2770 và kiểm chứng chuỗi đó trên APK gốc.')

heading('6 Tìm key và kiểm chứng trên APK gốc')
heading('6.1 Mở StringLiteral 2770 trong Listing',2)
para('Từ lời gọi Equals ở Hình 8, bấm đúp StringLiteral_2770 trong Decompiler. Listing chuyển đến vùng dữ liệu có nhãn này. Ảnh dưới giữ cả địa chỉ, nhãn và phần chú thích để đối chiếu cùng một vị trí.')
screenshot('a6686564-c88d-4cf5-850e-f92e0706a53d',(599,377,1497,881),'StringLiteral_2770 tại 044fce70 và nội dung BENTO-EASY-2025 trong Listing',16.8)
para('Trong Hình 10, vùng được chọn màu xanh nhạt mang nhãn StringLiteral_2770. Cột địa chỉ bên trái hiển thị 044fce70, còn chú thích bên phải hiển thị BENTO-EASY-2025. Đây chính là nhãn được truyền vào Equals ở Hình 8. Kết quả thu được ở bước này là chuỗi ứng viên BENTO-EASY-2025 tại địa chỉ Ghidra 0x044FCE70.')
para('Các chuỗi lân cận giúp phân biệt vai trò của dữ liệu. StringLiteral_2769 là BENTO SECURITY SYSTEM, StringLiteral_2771 là BENTO-SEC-, còn StringLiteral_2772 là BENTO_LICENSE_ACTIVATED. Chuỗi cuối dùng làm khóa lưu trạng thái kích hoạt. Chuỗi cần kiểm chứng là BENTO-EASY-2025 vì được tham chiếu trực tiếp trong phép so sánh của VerifyLicense.')
heading('6.2 Đối chiếu địa chỉ và giá trị trong stringliteral json',2)
para('Mở stringliteral.json thuộc cùng gói dump và tìm địa chỉ 43FCE70. Mục được tìm thấy phải chứa cả value và address, như ảnh dưới.')
oldfig(14,'Mục JSON có value BENTO-EASY-2025 và address 0x43FCE70',13)
para('Hình 11 cho thấy value là BENTO-EASY-2025, trùng chú thích tại StringLiteral_2770 trong Ghidra. address là 0x43FCE70, thấp hơn địa chỉ Listing một lượng 0x100000. Đây là Image Base đã kiểm tra ở Memory Map, nên hai địa chỉ được quy đổi như sau.')
code('Địa chỉ trong JSON       0x043FCE70\nImage Base của Ghidra    0x00100000\nĐịa chỉ trong Listing   0x044FCE70\n\n0x043FCE70 + 0x00100000 = 0x044FCE70')
para('Ghidra hiển thị 044fce70, còn JSON ghi 0x43FCE70; khác biệt này được giải thích bằng Image Base. Giá trị chuỗi và địa chỉ sau quy đổi đều khớp. Đối chiếu JSON xác nhận nguồn dữ liệu của nhãn, còn việc chuỗi có được chấp nhận hay không được kiểm tra bằng chạy ứng dụng ở bước sau.')
heading('6.3 Nhập key và kiểm tra kết quả trên APK gốc',2)
para('Mở APK gốc trên thiết bị Android, nhập BENTO-EASY-2025 rồi bấm XÁC NHẬN (VERIFY). Ảnh dưới ghi lại key đã nhập và thông báo sau khi xác nhận.')
fig(ROOT/'bento_report'/'evidence'/'original_apk_valid_key.png',(82,333,563,717),'APK gốc báo kích hoạt thành công sau khi nhập BENTO-EASY-2025',11.8)
para('Trong Hình 12, ô nhập giữ nguyên BENTO-EASY-2025 và phía dưới hiện thông báo màu xanh KÍCH HOẠT THÀNH CÔNG! ĐANG VÀO GAME... Kết quả này xác nhận chuỗi tìm được trong Listing và JSON được APK gốc chấp nhận.')
oldfig(16,'APK gốc chuyển vào màn chơi Level 1 sau khi xác nhận key',7.4)
para('Hình 13 cho thấy màn chơi Level 1 xuất hiện sau màn hình license. Hai ảnh ghi nhận hai trạng thái liên tiếp: ứng dụng báo chấp nhận key và trò chơi đã vào màn chơi. Đây là kết quả hoàn thành phần tìm key của bài thực hành.')

heading('7 Sửa VerifyLicense và xuất thư viện')
heading('7.1 Chuẩn bị bản sao và chọn lệnh cần sửa',2)
para('Trong Ghidra, lưu chương trình thành bản sao libil2cpp_patched để giữ phần phân tích của bản gốc. Mở lại VerifyLicense tại 0x1F7689C. Chọn lệnh đầu trong Listing rồi mở Patch Instruction; vùng nhập lệnh được viền đỏ như trong ảnh.')
screenshot('009fb381-433a-4000-900f-f0df02066f38',(647,516,1252,576),'Vùng Patch Instruction tại lệnh đầu VerifyLicense trước khi thay mã',16)
para('Hình 14 hiển thị lệnh str x30, [sp, #-0x20]! tại 01f7689c. Ngay bên dưới, tại 01f768a0, là lệnh stp x20, x19, [sp, #0x10]. Đây là hai lệnh đầu của hàm trước khi patch. Địa chỉ cần chọn là 0x1F7689C trong chương trình Ghidra; File Offset 0x1E7289C trong dump là vị trí trong file, không dùng thay cho địa chỉ Go To ở bước này.')
heading('7.2 Thay hai lệnh để trả về thành công ngay',2)
para('Tại 0x1F7689C, thay lệnh str bằng mov w0, #0x1. Tại 0x1F768A0, thay lệnh stp bằng ret. Sau khi xác nhận hai thao tác Patch Instruction, Listing hiển thị mã mới như ảnh dưới.')
screenshot('66cfd80d-de62-4ff1-af14-ce955325e492',(645,522,1115,565),'Hai lệnh ARM64 sau patch là mov w0, #0x1 và ret',15)
para('Trong Hình 15, địa chỉ 01f7689c có các byte 20 00 80 52 và lệnh mov w0, #0x1. Địa chỉ 01f768a0 có các byte C0 03 5F D6 và lệnh ret. Lệnh mov đặt giá trị trả về bằng 1; ret kết thúc hàm ngay sau đó. Hai lệnh đều dài 4 byte, nên phần sửa chiếm 8 byte đầu của VerifyLicense.')
table(['Địa chỉ Ghidra','Lệnh trước khi sửa','Lệnh sau khi sửa'],[
('0x1F7689C','str x30, [sp, #-0x20]!','mov w0, #0x1'),
('0x1F768A0','stp x20, x19, [sp, #0x10]','ret'),
],[4.6,7.0,5.3])
heading('7.3 Đọc lại Decompiler sau khi patch',2)
para('Sau khi sửa Listing, xem lại cùng hàm trong Decompiler để kiểm tra tác động của mã mới.')
screenshot('66cfd80d-de62-4ff1-af14-ce955325e492',(1246,170,1915,322),'Decompiler của VerifyLicense sau patch chỉ còn return 1',16.8)
para('Hình 16 hiển thị tên LicenseVerifier$$VerifyLicense và thân hàm return 1. So với Hình 8 và Hình 9, luồng hiện tại không còn đi qua IsNullOrWhiteSpace, Trim, Equals hoặc các nhánh kiểm tra bổ sung. Ghidra có thể đổi phần tham số hiển thị thành void vì mã mới không sử dụng đầu vào; tên nhãn vẫn xác định đúng hàm đã sửa.')
para('Mã trả về được đặt trước các lệnh cấp phát stack và gọi hàm con, nên hàm kết thúc mà chưa làm thay đổi stack theo phần thân cũ. Các lệnh còn lại vẫn có thể xuất hiện phía dưới trong Listing nhưng không được thực thi từ điểm vào này. Kết quả ở Ghidra là bằng chứng phần mã đã thay đổi; hiệu lực trên ứng dụng được kiểm tra sau khi xuất thư viện và build APK ở phần 8.')
heading('7.4 Xuất file ELF chứa byte đã sửa',2)
para('Lưu chương trình rồi mở File → Export Program. Chọn Format là Original File và đặt file đích trong C:\\Thuchanh2. Ảnh dưới ghi lại định dạng xuất trước khi xác nhận.')
screenshot('f7081120-b76f-4092-b76a-c3cda0952f70',(903,428,1219,628),'Chọn Original File và đường dẫn xuất thư viện đã sửa',10.2)
para('Hình 17 hiển thị Format Original File và Output File C:\\Thuchanh2\\libil2cpp_patched. Việc lưu chương trình trong project chỉ giữ cơ sở dữ liệu phân tích của Ghidra; bước Export tạo ra file thư viện để thay trong thư mục APK.')
screenshot('048ee7fd-98e5-43f2-ba4c-034ebc53bb5b',(898,423,1225,624),'Tích Export User Byte Modifications trong Options khi xuất',10.2)
para('Trong Hình 18, Export User Byte Modifications đã được tích. Tùy chọn này đưa các byte vừa patch vào file xuất. Sau khi xác nhận Options và Export, kiểm tra kết quả ở Export Results Summary.')
screenshot('8da9a840-6263-4f6e-978a-dc8c373ccf51',(700,750,1447,838),'Kết quả xuất file đích với định dạng Original File',15.3)
para('Hình 19 cho biết file đích là C:\\Thuchanh2\\libil2cpp_patched, Format là Original File và kích thước 72.828.928 byte. File xuất được đổi tên thành libil2cpp.so để khớp tên thư viện trong APK. Kích thước được giữ nguyên phù hợp với việc thay 8 byte tại chỗ, không chèn thêm lệnh vào file.')
heading('7.5 Thay thư viện trong thư mục APK đã Decompile',2)
para('Copy libil2cpp.so vừa xuất vào lib\\arm64-v8a của thư mục BentoLoop đã Decompile. Windows báo đã có file cùng tên, như trong ảnh dưới.')
screenshot('b9767d67-a577-45d4-a515-bbb923631f8e',(532,254,1316,847),'Xác nhận thay libil2cpp.so tại thư mục arm64-v8a',15.4)
para('Trong Hình 20, hộp thoại ghi rõ file trùng tên là libil2cpp.so và thư mục đích là arm64-v8a. Chọn Replace the file in the destination để thay thư viện cũ bằng bản đã patch. Sau khi copy hoàn tất, thư mục BentoLoop chứa thư viện mới được dùng làm đầu vào Compile ở phần 8. Phạm vi thay thư viện trong bài là bản ARM64.')

# Move the newly authored elements into the original section slot.
new_nodes=[node for node in body if node not in baseline and node.tag!=qn('w:sectPr')]
for node in new_nodes:anchor.addprevious(node)

# Refresh figure numbering for the retained later sections without duplicating images.
caption_count=0
for p in d.paragraphs:
    if re.match(r'^Hình \d+\.',p.text):
        caption_count+=1
        text=re.sub(r'^Hình \d+\.',f'Hình {caption_count}.',p.text)
        p.clear();r=p.add_run(text)
        r.bold=True;r.italic=True;r.font.size=Pt(10.5);r.font.color.rgb=RGBColor(0,0,0)
    if 'còn ảnh ở mục 6.2 thuộc APK gốc' in p.text:
        p.text=p.text.replace('còn ảnh ở mục 6.2 thuộc APK gốc','còn ảnh ở mục 6.3 thuộc APK gốc')

used={node.get(qn('r:embed')) for node in d.element.xpath('.//a:blip')}
for rid,rel in list(d.part.rels.items()):
    if rel.reltype.endswith('/image') and rid not in used:d.part.drop_rel(rid)
settings=d.settings.element
update=settings.find(qn('w:updateFields'))
if update is None:update=OxmlElement('w:updateFields');settings.append(update)
update.set(qn('w:val'),'true')
d.save(OUT)
assert hashlib.sha256(REF.read_bytes()).hexdigest()==ref_hash
check=Document(OUT)
alltext='\n'.join(check.element.xpath('.//w:t/text()'))
assert not re.search(r'\b(tôi|chúng tôi|chúng ta|mình|bạn)\b',alltext,re.I)
assert caption_count==26 and len(check.inline_shapes)==27
image_parts=[rel.target_part.blob for rel in check.part.rels.values() if rel.reltype.endswith('/image')]
assert len(set(hashlib.sha256(b).digest() for b in image_parts))==len(image_parts)
manifest={'reference':str(REF),'reference_sha256':ref_hash,'output':str(OUT),'new_figures':figures[7:],'figure_count':26,'changed_sections':['5','6','7'],'dependent_reference_update':'8.4 references 6.3','all_figure_images_unique':True,'render_pending':True}
(WORK/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print('OUTPUT',OUT)
print('FIGURES',caption_count,'SHAPES',len(check.inline_shapes))
