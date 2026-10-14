from pathlib import Path
from copy import deepcopy
import shutil, re, json
from docx import Document

ROOT = Path(__file__).resolve().parent
PATH = ROOT.parent / 'Bao cao dich nguoc SandSort bo sung day du.docx'
shutil.copy2(PATH, ROOT / 'before_personal_prose.docx')
doc = Document(PATH)
original_images = len(doc.inline_shapes)

def find(prefix):
    hits = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(hits) == 1, (prefix, len(hits))
    return hits[0]

def replace(prefix, text):
    p = find(prefix)
    props = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    r = p.add_run(text)
    if props is not None:
        r._r.insert(0, props)

def delete(prefix):
    p = find(prefix)
    assert not p._p.xpath('.//w:drawing')
    p._p.getparent().remove(p._p)

changes = [
('Bài thực hành khảo sát', 'Bài thực hành phân tích gói game SandSort bằng UABEA, Il2CppDumper, dnSpy và Ghidra. Nội dung gồm khảo sát tài nguyên hình ảnh, tìm nơi đọc và lưu số coin khi vào game, sau đó lần theo luồng nhận thưởng để xác định cách cập nhật coin. Các kết quả được rút ra từ phân tích tĩnh, chưa được kiểm chứng khi game chạy.'),
('Thực hành được thực hiện', r'Môi trường thực hành là Windows trong máy ảo VMware Workstation có tên thuchanhdichnguoc1. Dữ liệu và công cụ được đặt trong C:\LabGame. File mẫu SandSort_14-9.zip nằm trong APK_goc, còn dữ liệu đã giải nén nằm trong APK_giai_nen. Thư viện được chọn để phân tích là bản ARM64.'),
('Gói mẫu được lưu tại', r'File mẫu được lưu tại C:\LabGame\APK_goc\SandSort_14-9.zip, có dung lượng hiển thị là 110 605 KB. Khi mở file, các thư mục assets, lib, res cùng AndroidManifest.xml và các file classes.dex xuất hiện ở thư mục gốc. Cấu trúc này phù hợp với một gói ứng dụng Android. Tuy nhiên, ảnh hiện có chỉ xác nhận tên file ZIP, chưa xác định tên APK ban đầu hoặc phiên bản game.'),
('Thư mục C:\\LabGame tập hợp', r'Trong C:\LabGame, APK_goc chứa file mẫu, APK_giai_nen chứa dữ liệu đã giải nén và Asset_xuat chứa tài nguyên đã xuất. Các thư mục còn lại chứa công cụ phân tích. Kết quả từ Il2CppDumper được lưu trong Il2CppDumper-win-v6.7.46 (1).'),
('Project Ghidra có tên', r'Project Ghidra được đặt tên là dichnguoc. Các file dump.cs, script.json và thư mục DummyDll được lưu trong C:\LabGame\Il2CppDumper-win-v6.7.46 (1) để tra cứu và gắn tên phương thức.'),
('Trong C:\\LabGame\\uabea-windows', r'UABEANext4.Desktop.exe được mở từ C:\LabGame\uabea-windows. Cửa sổ hiển thị Workspace Explorer và Inspector. Công cụ này được dùng để xem và xuất asset ở phần 2.'),
('Trong C:\\LabGame\\dnSpy-net-win64', r'dnSpy.exe được mở từ C:\LabGame\dnSpy-net-win64. Thanh tiêu đề hiển thị phiên bản 6.1.8, 64 bit .NET. Assembly Explorer được dùng để tra cứu các lớp trong DummyDll ở phần 4.'),
('Chạy Il2CppDumper.exe trong', r'Il2CppDumper.exe được mở từ C:\LabGame\Il2CppDumper-win-v6.7.46 (1). Chương trình xuất hiện hộp chọn file binary như Hình 6. Ở bước này cần chọn libil2cpp.so trước, sau đó chọn global-metadata.dat trong hộp tiếp theo.'),
('Chạy ghidraRun.bat trong', r'Ghidra được mở bằng ghidraRun.bat trong C:\LabGame\ghidra_12.1.2_PUBLIC_20260605\ghidra_12.1.2_PUBLIC. Cửa sổ quản lý hiển thị project dichnguoc và dòng Ghidra startup complete. Các bước import và phân tích thư viện được trình bày ở phần 5.'),
('Thư viện native sử dụng xuyên suốt', r'Thư viện được phân tích là C:\LabGame\APK_giai_nen\lib\arm64-v8a\libil2cpp.so. Import Results Summary của Ghidra hiển thị định dạng ELF shared object, processor AARCH64 và các mã nhận diện bên dưới. Các mã băm này thuộc libil2cpp.so, không phải toàn bộ gói APK.'),
('Dữ liệu gói ứng dụng được giải nén', r'Dữ liệu được giải nén vào C:\LabGame\APK_giai_nen. Thư mục này có assets, lib, res, các file classes.dex, classes2.dex và AndroidManifest. assets chứa dữ liệu Unity, còn lib chứa thư viện native theo kiến trúc. Các file DEX và manifest thuộc phần ứng dụng Android; mã IL2CPP cần phân tích nằm trong thư viện native.'),
('File Unity được mở để khảo sát', r'File được mở bằng UABEA là C:\LabGame\APK_giai_nen\assets\bin\Data\data.unity3d. Trong UABEA, sharedassets0.assets, sharedassets1.assets và sharedassets2.assets được hiển thị bên trong workspace data.unity3d. sharedassets0.assets được chọn để tìm tài nguyên hình ảnh.'),
('File Explorer bên dưới xác nhận', 'Hình 10 cho thấy vị trí file data.unity3d trong thư mục đã giải nén. Nội dung của file được mở bằng UABEA ở mục tiếp theo.'),
('Trong workspace data.unity3d, chọn', 'Trong workspace data.unity3d, asset bottle 1 được tìm thấy ở mục sharedassets0.assets và có loại Texture2D. Preview hiển thị một chai chứa vật liệu màu vàng cam. Kích thước texture là 281 × 552 pixel.'),
('Theo hình dạng trong Preview', 'Hình ảnh trong Preview cho thấy texture này có thể được dùng để hiển thị chai chứa cát hoặc vật liệu màu trong game. Tuy nhiên, prefab và scene chưa được kiểm tra để xác định chính xác vị trí sử dụng.'),
('Sau khi xem Preview, sử dụng', r'Sau khi xem Preview, texture được xuất bằng chức năng Export của UABEA vào C:\LabGame\Asset_xuat. File xuất có tên bottle_1_Texture2D_sharedassets0.assets.png và đuôi .png. Hình thu nhỏ của file giống asset bottle 1 trong Preview.'),
('Il2CppDumper được sử dụng để đối chiếu', 'Il2CppDumper kết hợp libil2cpp.so với global-metadata.dat để tạo dữ liệu tra cứu lớp, phương thức và địa chỉ mã máy. Hai file lấy từ cùng gói đã giải nén. Bản libil2cpp.so ARM64 cũng được dùng khi import vào Ghidra.'),
('Đầu vào binary:', r'Đầu vào binary: C:\LabGame\APK_giai_nen\lib\arm64-v8a\libil2cpp.so. Đây là bản ARM64 được dùng để phân tích trong Ghidra.'),
('Hai ảnh File Explorer xác nhận', 'Hình 13 và Hình 14 cho thấy vị trí hai file đầu vào. Khi chạy Il2CppDumper, libil2cpp.so được chọn trước, sau đó là global-metadata.dat. Hai ảnh này được chụp trong File Explorer, không phải trong hộp chọn file của lần chạy dump.'),
('Thư mục kết quả Il2CppDumper', 'Thư mục Il2CppDumper-win-v6.7.46 (1) có DummyDll, dump.cs, script.json, ghidra.py, il2cpp.h và stringliteral.json. Các file này được dùng để tra cứu tên, cấu trúc và địa chỉ; chúng không phải toàn bộ mã nguồn C# gốc của game.'),
('Quy trình chuyển từ tên phương thức', 'Tên lớp và phương thức được tìm trong DummyDll bằng dnSpy, sau đó tra phương thức tương ứng trong dump.cs để lấy RVA và field offset. Tiếp theo, ghidra.py đọc script.json để gắn tên trong Ghidra. Từ tên đã gắn hoặc địa chỉ tính bằng RVA + Image Base, có thể mở Listing và Decompiler để đọc mã native. DummyDll chỉ cung cấp thông tin phục vụ tra cứu, không chứa phần thực thi native của phương thức.'),
('Mở DummyDll\\Assembly-CSharp.dll', r'Khi mở DummyDll\Assembly-CSharp.dll trong dnSpy và tìm từ khóa coin ở namespace FoodSort, PlayerResourceManager có field private int _coin với FieldOffset 0x20. Kiểu int phù hợp với số lượng coin. Các lệnh đọc và ghi field này được kiểm tra tiếp trong mã native ở phần 6.'),
('Kết quả tìm kiếm cũng có', 'Kết quả tìm kiếm còn có _coinGain trong UIWinGame. Tuy nhiên, field này có kiểu Transform nên được loại khỏi ứng viên chứa số dư coin. Field int _coin của PlayerResourceManager được giữ lại để kiểm tra.'),
('Các phương thức get_Coin, set_Coin', 'Các phương thức được kiểm tra gồm get_Coin, set_Coin, Awake, Start và StateCoin. Địa chỉ của get_Coin và set_Coin được tra trong thuộc tính Address và dump.cs. Với Awake, Start và StateCoin, địa chỉ được lấy từ Listing ở phần 6 và 7; RVA được tính bằng cách trừ Image Base 0x100000. Token metadata và RVA của mã IL trong DummyDll không phải địa chỉ mã native.'),
('get_Coin trả về int, không có', 'get_Coin trả về int và không có tham số. set_Coin nhận int value, trả về void. Thân hàm rỗng hoặc return 0 trong DummyDll chỉ là phần được dựng để tra cứu; cần đọc mã native để biết hành vi thực tế.'),
('Trong ảnh dump.cs, File Offset', 'Hình 19 ghi File Offset của get_Coin là 0x229C708 và set_Coin là 0x229C710. Đây là vị trí trong file ELF. Với các phương thức còn lại, địa chỉ được đối chiếu trong Listing và RVA được tính bằng cách trừ Image Base 0x100000.'),
('Tạo project Non-Shared tên', 'Project Non-Shared được tạo với tên dichnguoc và import file libil2cpp.so đã dùng với Il2CppDumper. Hộp import nhận dạng định dạng ELF và language AARCH64:LE:64:v8A:default. Sau khi mở CodeBrowser, bước phân tích tự động được chạy để nhận diện vùng mã, hàm và tham chiếu.'),
('Ghidra sử dụng Image Base', 'Image Base trong lần import này là 0x100000. Ví dụ, get_Coin có RVA 0x22A0708 và địa chỉ Ghidra 0x23A0708. File Offset chỉ vị trí dữ liệu trong file ELF, còn địa chỉ Ghidra chỉ vị trí trong chương trình đã import.'),
('Listing cung cấp lệnh ARM64', 'Listing hiển thị lệnh ARM64 và địa chỉ. Decompiler dựng mã giả để dễ đọc luồng xử lý, nhưng vẫn cần đối chiếu với khai báo dnSpy và các offset. Mã giả này không phải mã nguồn C++ hoặc C# gốc.'),
('Thêm thư mục Il2CppDumper', 'Thư mục Il2CppDumper-win-v6.7.46 (1) được thêm vào Script Directories. Lần chạy đầu xuất hiện thông báo Ghidra was not started with PyGhidra. Python is not available, nên script chưa chạy được.'),
('Trong cửa sổ quản lý project, mở', 'Để xử lý lỗi, Jython được bật trong File > Install Extensions, sau đó khởi động lại Ghidra. Script Il2CppDumper sử dụng Python 2/Jython, nên chỉ thị runtime được thêm vào đầu ghidra.py như sau:'),
('Chạy lại ghidra.py và chọn', 'Sau khi chạy lại ghidra.py với script.json của cùng lần dump, Console hiển thị Script finished! và ghidra.py> Finished!. Tên các phương thức FoodSort xuất hiện trong Symbol Tree và project được lưu lại.'),
('Kết quả này xác định vị trí field', 'Hai lệnh xác nhận field _coin nằm ở offset 0x20 trong đối tượng PlayerResourceManager. 0x20 là độ lệch của field, không phải lượng coin hoặc địa chỉ cố định khi game chạy. Decompiler hiển thị undefined4 cho dữ liệu 4 byte, còn dnSpy xác định kiểu của field là int.'),
('Hai hàm accessor không chứa', 'get_Coin và set_Coin chỉ đọc hoặc ghi field _coin. Trong đoạn mã đã xem, hai hàm không cộng coin và không gọi PlayerPrefs. Start được kiểm tra tiếp để tìm bước nạp số dư, còn StateCoin được kiểm tra để tìm bước cập nhật.'),
('Awake tại 0x23A15D0', 'Awake tại 0x23A15D0 chỉ có lệnh ret, nên được loại khỏi ứng viên trực tiếp nạp hoặc khởi tạo coin trong bản thư viện này.'),
('Luồng giá trị là dữ liệu đã lưu', 'Start lấy số dư đã lưu bằng GetInt rồi ghi vào _coin. Khi khóa đã tồn tại, số coin được lấy từ lần lưu trước; khi chưa có khóa, giá trị mặc định là 100. Vì vậy, 100 không phải số coin được cộng lại mỗi lần vào game. Các lời gọi GetInt tại offset 0x24, 0x28 và 0x2C thuộc các field booster.'),
('Đi theo _StringLiteral_3130', 'Khi mở _StringLiteral_3130 trong Listing, nhãn có chú thích Coin_SandSort. Start dùng khóa này để đọc bằng GetInt, còn StateCoin dùng cùng khóa để lưu bằng SetInt.'),
('Mã diễn giải tương đương', 'Phần nạp coin có thể diễn giải bằng đoạn mã dưới đây. Đây là cách viết lại hành vi đã quan sát, không phải mã nguồn gốc:'),
('Kết luận: Start thực sự', 'Start đọc số coin từ khóa Coin_SandSort và ghi vào field _coin. Khi chưa có dữ liệu lưu, hàm dùng giá trị mặc định 100.'),
('Phần tiếp theo phân biệt', 'StateCoin thực hiện việc cộng hoặc trừ số dư theo tham số truyền vào. Trong nhánh NextLevel đã kiểm tra, hàm được gọi với state = true và coin = 100.'),
('Listing bên dưới xác nhận', 'Hình 27 hiển thị điểm bắt đầu StateCoin tại 0x23A16F8. Hình 28 cho thấy phép tính và bước lưu số dư.'),
('Với coin là số dương', 'Với coin là số dương, state = true làm tăng số dư, còn state = false làm giảm số dư. Phía sau có các lời gọi RecordLevelCoinChange và TrackWhenMetadataReady. Phần xử lý này chưa được phân tích đầy đủ nên chưa xác định có bước xác thực số dư trên máy chủ hay không.'),
('UIWinGame.NextLevel tại', 'UIWinGame.NextLevel tại 0x23B6044, RVA 0x22B6044, chuẩn bị state machine <NextLevel>d__26 và gọi AsyncVoidMethodBuilder.Start. Phần xử lý nhận thưởng nằm trong <NextLevel>d__26.MoveNext. Trong nhánh này, TryStartClaim phải trả về true trước khi lấy PlayerResourceManager và gọi StateCoin.'),
('Dòng 42 - 43 của ảnh dưới', 'Dòng 42 - 43 trong Hình 30 truyền state = 1 và coin = 100. Listing có lệnh mov w1, #0x1 và mov w2, #0x64; 0x64 tương đương 100. 0x23B7804 là địa chỉ lệnh gọi StateCoin, không phải địa chỉ bắt đầu NextLevel hoặc MoveNext.'),
('Hai tham số chuỗi ở lời gọi', 'Hai tham số chuỗi tại lời gọi có nhãn _StringLiteral_13854 và _StringLiteral_14787. Nội dung chuỗi chưa được xác định nên chưa biết tên nguồn thưởng. Giá trị 100 ở lời gọi này là lượng coin được cộng, còn 100 trong Start là giá trị mặc định khi chưa có dữ liệu lưu.'),
('Sau khi bắt đầu nhận thưởng', 'Sau khi bắt đầu nhận thưởng, dòng 24–26 ghi 0 vào field +0x20 của UIClick tại this +0x50. Dòng 33–36 làm tương tự với UIClick tại this +0x58 nếu đối tượng tồn tại. Đối chiếu mục 7.5 cho thấy đây là thao tác xóa _action của nút tiếp tục và nút X2. Việc xóa _action khiến các nút không còn callback để tiếp tục kích hoạt trong lúc nhận thưởng.'),
('Như vậy, cờ ban đầu false', 'Nếu _isClaimingReward ban đầu là false, TryStartClaim cho phép luồng nhận thưởng tiếp tục. Khi cờ đã là true, hàm trả về false và nhánh cộng coin bị bỏ qua. Vì vậy, hai lần gọi liên tiếp khi cờ chưa được đặt lại không đều nhận thưởng. Kết quả này áp dụng khi hàm hoàn tất bình thường; trường hợp gặp đối tượng null có thể đi vào nhánh xử lý lỗi.'),
('Trong dnSpy, field private bool', 'Trong dnSpy, _isClaimingReward có kiểu bool và FieldOffset 0xA8, khớp vị trí được đọc trong TryStartClaim. Field _isRewardAdCallbackHandled nằm ở offset 0xA9 và là một cờ riêng.'),
('UIWinGame.OnEnable có RVA', 'UIWinGame.OnEnable có RVA 0x22B5B40, địa chỉ Ghidra 0x23B5B40. Khi UIManager tồn tại, dòng 71 ghi 0 bằng kiểu undefined2 tại +0xA8, đặt lại cả _isClaimingReward (+0xA8) và _isRewardAdCallbackHandled (+0xA9). Dòng 72 gọi BindClaimActions để gắn lại callback. Phần đã kiểm tra chỉ cho thấy cách đặt lại cờ trong OnEnable, chưa xác định số lần giao diện này được mở trong một level.'),
('Trong FoodSort.UIClick, field', 'Trong FoodSort.UIClick, field private Action _action có FieldOffset 0x20. Phép ghi tại UIClick +0x20 vì vậy là cập nhật callback của nút.'),
('FoodSort.UIWinGame.BindClaimActions bắt đầu', 'FoodSort.UIWinGame.BindClaimActions bắt đầu tại 0x23B5D30, tương ứng RVA 0x22B5D30 với Image Base 0x100000. Dòng 22 lấy UIClick tại this +0x50. Dòng 24 tạo System.Action gắn đối tượng UIWinGame hiện tại với NextLevel, rồi dòng 27 ghi Action vào UIClick +0x20. Ở bước này, NextLevel được gắn làm callback của nút và sẽ chạy khi callback được kích hoạt.'),
('Với đối tượng ở this +0x58', 'Với UIClick tại this +0x58, hàm kiểm tra đối tượng UnityEngine.Object khác null rồi tạo System.Action gắn với ClaimX2Reward ở dòng 39. Action được ghi vào field _action tại +0x20. ClaimX2Reward được xác định là callback của nút X2, nhưng phần xử lý quảng cáo và lượng coin được cộng chưa được phân tích.'),
('Lời gọi trong OnPointerUp', 'OnPointerUp gọi callback qua delegate nên tại dòng gọi không xuất hiện trực tiếp tên NextLevel. BindClaimActions cho thấy NextLevel đã được gắn cho _uIClickNextLevel. Kết quả này áp dụng cho nút đó; các đối tượng UIClick khác có thể được gắn callback khác.'),
('Chuỗi xử lý đã xác định là:', 'Khi UIWinGame được bật, OnEnable đặt lại hai cờ và gọi BindClaimActions để gắn callback cho các nút. Khi nhả con trỏ trong vùng nút tiếp tục, OnPointerUp gọi NextLevel. Phần MoveNext kiểm tra TryStartClaim; nếu được phép, nó gọi StateCoin với state = true và coin = 100. StateCoin cộng vào _coin rồi lưu số dư mới bằng PlayerPrefs.SetInt và PlayerPrefs.Save.'),
('NextLevel không chứa trực tiếp', 'NextLevel chỉ khởi chạy state machine, nên cần đọc MoveNext để tìm phần xử lý nhận thưởng.'),
('Đã xác định đối tượng _uIClickNextLevel', 'Luồng nhận thưởng đã được lần từ nút tiếp tục đến thao tác cộng và lưu coin. Tuy nhiên, chưa xác định hàm xử lý kết quả thắng nào mở UIWinGame, cũng chưa kiểm chứng luồng này khi game chạy. Do đó, chưa thể khẳng định mọi lần thắng đều tự động cộng ngay 100 coin. Nhánh X2 mới được xác định ở bước gắn callback.'),
('Nhận định có cơ sở từ ba bằng chứng', 'Start đọc số dư từ PlayerPrefs, còn StateCoin cập nhật và lưu số dư ngay trên client. Vì vậy, việc sửa dữ liệu cục bộ hoặc can thiệp bước cập nhật có thể làm thay đổi số coin. Chưa có thử nghiệm can thiệp để kiểm tra tác động thực tế.'),
('Tác động dự kiến là số dư', 'Nếu số dư phía client bị thay đổi, phần hiển thị và các thao tác tiêu coin có thể bị ảnh hưởng khi chúng sử dụng số dư này. Các ảnh hiện có chưa cho thấy máy chủ có chấp nhận số dư đã sửa hay số coin đó có dùng được cho mọi chức năng hay không.'),
('Phân tích tĩnh chưa xác minh', 'Dữ liệu PlayerPrefs khi game chạy và số coin hiển thị sau khi thắng chưa được kiểm tra. Các nhánh thưởng khác, nội dung hai chuỗi mô tả nguồn thưởng và các hàm thống kê trong StateCoin cũng chưa được phân tích đầy đủ. Do đó, kết quả hiện tại chưa đủ để khẳng định có thể khai thác thành công hoặc game không có cơ chế bảo vệ.'),
('Với game có tài khoản và máy chủ', 'Với game có máy chủ, có thể giảm phụ thuộc vào client bằng cách lưu số dư và kiểm tra điều kiện nhận thưởng trên máy chủ. Client gửi yêu cầu, còn máy chủ xác định lượng coin hợp lệ và trả về số dư mới. Mỗi lượt thưởng cần có mã riêng để tránh xử lý cùng một yêu cầu nhiều lần. Dữ liệu lưu trên client chỉ dùng để hiển thị tạm và được cập nhật theo kết quả từ máy chủ.'),
('Trong chế độ offline, kiểm tra', 'Với game offline, kiểm tra tính toàn vẹn có thể làm việc sửa dữ liệu khó hơn, nhưng dữ liệu vẫn nằm trên thiết bị của người chơi. Cờ _isClaimingReward giúp hạn chế việc nhận thưởng lặp trong luồng giao diện đã xem; cờ này không cho thấy số dư đã được máy chủ xác thực.'),
('Bài thực hành đã xác định', 'Qua bài thực hành, texture bottle 1 trong sharedassets0.assets đã được xuất thành file PNG và nơi đọc, cập nhật số coin đã được tìm thấy trong mã native. Start đọc số dư từ khóa Coin_SandSort, dùng mặc định 100 khi chưa có dữ liệu lưu. Với nút tiếp tục của UIWinGame, callback đi từ BindClaimActions qua OnPointerUp đến NextLevel; nhánh MoveNext gọi StateCoin để cộng 100 coin và lưu số dư mới. Bước xử lý kết quả thắng mở UIWinGame và hành vi khi game chạy vẫn cần được kiểm chứng.'),
]

for prefix, text in changes:
    replace(prefix, text)
for prefix in [
    'Kết quả phân tích tĩnh cho thấy',
    'Các ảnh ở mục 1.1 đến 1.3',
    'Gói đầu vào đang lưu có tên',
    'Kết quả gồm đường dẫn đầy đủ',
]:
    delete(prefix)

# Keep the existing figures, tables, evidence, styles and cover intact.
assert len(doc.inline_shapes) == original_images == 41
captions = [p.text for p in doc.paragraphs if re.match(r'^Hình \d+\.', p.text)]
assert len(captions) == 40
body_text = '\n'.join(p.text for p in doc.paragraphs)
assert not re.search(r'\b(nhóm|chúng tôi|chúng em|tôi|em)\b', body_text, re.I)
for key in ['0x23B5D30','0x23B11BC','0x23B7804','0x23A15D4','0x23A16F8',
            'Coin_SandSort','bottle_1_Texture2D_sharedassets0.assets.png',
            '0xA8','0xA9','global-metadata.dat','281 × 552']:
    assert key in body_text, key
doc.save(PATH)
(ROOT / 'personal_prose_summary.json').write_text(json.dumps({
    'replaced_paragraphs':len(changes),'deleted_paragraphs':4,
    'figures':len(captions),'inline_shapes':original_images,
    'output':str(PATH)
}, ensure_ascii=False, indent=2),encoding='utf-8')
print('PROSE_UPDATED', len(changes), 'paragraphs; 40 figures preserved')
