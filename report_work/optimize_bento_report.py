from pathlib import Path
import hashlib, json, re
from PIL import Image
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
WORK = ROOT / 'bento_report_optimized'
WORK.mkdir(exist_ok=True)
REF = ROOT.parent / 'Báo cáo BentoLoop làm rõ phần 5 6 7.docx'
OUT = ROOT.parent / 'Báo cáo BentoLoop tối ưu nội dung.docx'
TEMP = Path(r'C:\Users\ADMIN\AppData\Local\Temp')
reference_hash = hashlib.sha256(REF.read_bytes()).hexdigest()
d = Document(REF)
cover = d.tables[0]._tbl
first = d.element.body[0]
for element in list(d.element.body):
    if element is not first and element is not cover and element.tag != qn('w:sectPr'):
        d.element.body.remove(element)
original = (ROOT / 'build_bento_report.py').read_text(encoding='utf-8-sig')
helpers = original[original.index('def para(text):'):original.index('d.add_page_break()')]
exec(compile(helpers, str(ROOT / 'build_bento_report.py'), 'exec'))

def retained(folder, number, caption, width=16, box=None):
    path = ROOT / folder / f'figure-{number:02d}.png'
    with Image.open(path) as im:
        bounds = box or (0, 0, im.width, im.height)
    return fig(path, bounds, caption, width)

d.add_page_break()
p = d.add_paragraph('MỤC LỤC')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Pt(0)
for r in p.runs:
    r.bold = True
    r.font.size = Pt(18)
p = d.add_paragraph()
p.paragraph_format.first_line_indent = Pt(0)
for kind in ['begin', 'instruction', 'separate', 'end']:
    run = p.add_run()
    if kind == 'instruction':
        el = OxmlElement('w:instrText')
        el.set(qn('xml:space'), 'preserve')
        el.text = ' TOC \\o "1-2" \\h \\z \\u '
    else:
        el = OxmlElement('w:fldChar')
        el.set(qn('w:fldCharType'), kind)
    run._r.append(el)
d.add_page_break()

heading('1 Mục tiêu và môi trường thực hành')
para('Bài thực hành phân tích cơ chế kiểm tra license của BentoLoop, tìm key hợp lệ và sửa hàm kiểm tra để build lại APK. Kết quả thu được là key BENTO-EASY-2025 được APK gốc chấp nhận; bản APK đã sửa báo kích hoạt thành công với chuỗi vinh123.')
para('Quá trình phân tích thực hiện trên Windows trong máy ảo VMware. APK Easy Tool được dùng để Decompile, Compile và ký APK; dữ liệu dump do giảng viên cung cấp giúp tra cứu lớp và địa chỉ; Ghidra dùng để đọc và sửa mã native. Kết quả được kiểm thử trên thiết bị Android kết nối với máy thật.')

heading('2 Chuẩn bị thư viện và xác định VerifyLicense')
heading('2.1 Decompile bằng APK Easy Tool', 2)
para('Trong APK Easy Tool, chọn BentoLoop.apk bằng Browse rồi bấm Decompile. Công cụ tạo thư mục BentoLoop chứa các thành phần của ứng dụng. Đây là thư mục làm việc dùng cả khi lấy thư viện để phân tích và khi Compile APK sau khi sửa.')
retained('bento_report', 1, 'Chọn BentoLoop.apk và sử dụng Decompile trong APK Easy Tool', 13.8, (0, 0, 692, 477))
para('Thư viện cần phân tích là libil2cpp.so trong lib\\arm64-v8a của thư mục đã Decompile. Đây là bản ARM64 chứa mã native của ứng dụng Unity IL2CPP. Thư viện này được mở trong Ghidra và cũng là file cần thay trước khi build lại APK.')
heading('2.2 Tra cứu hàm trong dữ liệu dump', 2)
para('Trong dump.cs, tìm LicenseVerifier thuộc namespace GeckoArrow.Runtime.LoadingSystem. Khai báo VerifyLicense(string inputKey) cho biết hàm nhận chuỗi nhập vào và trả về bool. Thân hàm rỗng trong dump chỉ phục vụ tra cứu; logic thực thi được đọc trong libil2cpp.so.')
retained('bento_report', 4, 'Khai báo LicenseVerifier và RVA của VerifyLicense trong dump.cs', 16, (0, 0, 664, 241))
para('RVA của VerifyLicense là 0x1E7689C. Chương trình Ghidra có Image Base 0x100000, nên địa chỉ cần mở bằng Go To là 0x1F7689C. Tên hàm và nhãn chuỗi được gắn từ script.json của cùng gói dump để thuận tiện đối chiếu mã native.')
code('Địa chỉ Ghidra = Image Base + RVA\n0x00100000 + 0x01E7689C = 0x01F7689C')

heading('3 Phân tích và kiểm chứng license key')
heading('3.1 Theo dõi phép so sánh và giá trị trả về', 2)
para('Mở VerifyLicense tại 0x1F7689C và theo dõi đoạn xử lý đầu vào trong Decompiler. Ba lời gọi IsNullOrWhiteSpace, Trim và Equals dẫn đến đầu mối tìm key.')
f_input = retained('bento_report_expanded', 8, 'VerifyLicense xử lý đầu vào và so sánh với StringLiteral_2770', 16.8)
para(f'Trong Hình {f_input}, dòng 19 kiểm tra chuỗi rỗng hoặc chỉ có khoảng trắng. Khi kết quả bằng 0, luồng tiếp tục tới dòng 25 để Trim khoảng trắng ở hai đầu. Dòng 26 gọi String.Equals(uVar3, StringLiteral_2770, 4, 0), so sánh chuỗi đã xử lý với dữ liệu mang nhãn StringLiteral_2770. Nhãn này biểu diễn một chuỗi, không phải tên hàm.')
f_return = retained('bento_report_expanded', 9, 'Các nhánh cuối VerifyLicense trả về 1 khi thành công và 0 khi thất bại', 16.8)
para(f'Hình {f_return} cho thấy nhánh thành công gán uVar3 = 1 ở dòng 51, sau đó return uVar3 ở dòng 57. Khi phép so sánh trực tiếp khớp, luồng đi tới kết quả này. Nếu không khớp, hàm còn thử các nhánh kiểm tra bổ sung; khi đều thất bại, luồng chuyển tới LAB_01f76914 và gán uVar3 = 0. Vì vậy, cần đọc nội dung StringLiteral_2770 để tìm chuỗi được nhánh trực tiếp chấp nhận.')
heading('3.2 Đọc StringLiteral 2770 và đối chiếu địa chỉ', 2)
para('Bấm đúp StringLiteral_2770 trong lời gọi Equals để chuyển đến dữ liệu trong Listing. Vùng được chọn dưới đây hiển thị đồng thời nhãn chuỗi, địa chỉ và nội dung key.')
f_literal = retained('bento_report_expanded', 10, 'StringLiteral_2770 tại 044fce70 chứa BENTO-EASY-2025', 16.8, (0, 107, 898, 236))
para(f'Trong Hình {f_literal}, cột địa chỉ bên trái ghi 044fce70, nhãn là StringLiteral_2770 và chú thích bên phải là BENTO-EASY-2025. Đây chính là nhãn được truyền vào Equals ở Hình {f_input}, nên chuỗi này được chọn để kiểm chứng trên APK gốc.')
f_json = retained('bento_report_expanded', 11, 'stringliteral.json ghi cùng key với address 0x43FCE70', 13)
para(f'Hình {f_json} là mục tương ứng trong stringliteral.json: value bằng BENTO-EASY-2025 và address bằng 0x43FCE70. Địa chỉ JSON cộng Image Base 0x100000 được 0x044FCE70, khớp vị trí trong Listing. Hai nguồn dữ liệu vì thế cùng xác nhận một chuỗi và một vị trí sau quy đổi.')
code('0x043FCE70 + 0x00100000 = 0x044FCE70')
heading('3.3 Kiểm chứng trên APK gốc', 2)
para('Trên APK gốc, nhập BENTO-EASY-2025 rồi bấm XÁC NHẬN (VERIFY). Ảnh dưới ghi lại chuỗi đã nhập cùng thông báo KÍCH HOẠT THÀNH CÔNG! ĐANG VÀO GAME...')
retained('bento_report_expanded', 12, 'APK gốc chấp nhận key BENTO-EASY-2025', 10.8)
retained('bento_report_expanded', 13, 'APK gốc chuyển vào Level 1 sau khi xác nhận key', 7.0)
para('Thông báo kích hoạt và màn chơi Level 1 xác nhận key tìm được bằng phân tích tĩnh có hiệu lực trên APK gốc. Kết quả này hoàn thành phần tìm key và là mốc đối chiếu với bản APK sửa ở phần 5.')

heading('4 Sửa hàm và thay thư viện')
heading('4.1 Thay hai lệnh đầu VerifyLicense', 2)
para('Lưu chương trình Ghidra thành bản sao libil2cpp_patched rồi mở VerifyLicense tại 0x1F7689C. Hai lệnh đầu của hàm được chọn để sửa bằng Patch Instruction.')
f_before = retained('bento_report_expanded', 14, 'Hai lệnh đầu VerifyLicense trước khi sửa', 16)
para(f'Hình {f_before} ghi lại lệnh str x30, [sp, #-0x20]! tại 0x1F7689C và stp x20, x19, [sp, #0x10] tại 0x1F768A0. Đây là các lệnh chuẩn bị stack trước khi thực hiện phần kiểm tra license.')
f_after = retained('bento_report_expanded', 15, 'Hai lệnh sau khi sửa là mov w0, #0x1 và ret', 15)
para(f'Tại 0x1F7689C, thay str bằng mov w0, #0x1; tại 0x1F768A0, thay stp bằng ret. Hình {f_after} hiển thị các byte mới lần lượt là 20 00 80 52 và C0 03 5F D6. Lệnh mov đặt giá trị trả về bằng 1, còn ret kết thúc hàm ngay. Hai lệnh dài tổng cộng 8 byte.')
f_decomp = retained('bento_report_expanded', 16, 'Decompiler của VerifyLicense sau sửa chỉ còn return 1', 16.8)
para(f'Hình {f_decomp} cho thấy thân VerifyLicense đã rút gọn thành return 1. Từ điểm vào hàm, mã mới trả về thành công trước khi đi tới Trim, Equals hoặc các nhánh kiểm tra bổ sung. Hai lệnh cấp phát stack ban đầu đã được thay, nên việc trả về sớm không cần khôi phục stack do phần thân cũ tạo ra.')
heading('4.2 Xuất thư viện có chứa byte đã sửa', 2)
para('Mở File → Export Program, chọn Format là Original File và đặt file đích C:\\Thuchanh2\\libil2cpp_patched. Bước Export tạo file thư viện để đưa vào APK; thao tác Save trong project chỉ lưu cơ sở dữ liệu phân tích của Ghidra.')
retained('bento_report_expanded', 17, 'Xuất chương trình với định dạng Original File', 9.3)
retained('bento_report_expanded', 18, 'Bật Export User Byte Modifications để đưa byte đã sửa vào file xuất', 9.3)
para('Trong Options, tích Export User Byte Modifications rồi xác nhận xuất file. Tùy chọn này giữ phần byte đã patch trong kết quả. File xuất được đổi tên thành libil2cpp.so để khớp tên thư viện mà ứng dụng sử dụng.')
heading('4.3 Thay thư viện trong thư mục đã Decompile', 2)
para('Copy libil2cpp.so đã xuất vào lib\\arm64-v8a của thư mục BentoLoop đã Decompile. Khi Windows báo trùng tên, chọn Replace the file in the destination để thay đúng thư viện cũ.')
retained('bento_report_expanded', 20, 'Thay libil2cpp.so trong thư mục arm64-v8a', 14.5, (101, 0, 754, 295))
para('Sau bước này, thư mục BentoLoop chứa thư viện ARM64 đã sửa và được dùng làm đầu vào Compile. Địa chỉ patch và các lệnh trong báo cáo áp dụng cho bản thư viện đã phân tích.')

heading('5 Build APK và kiểm thử bản đã sửa')
heading('5.1 Compile và ký bằng APK Easy Tool', 2)
para('Trong APK Easy Tool, dùng Select decompiled APK để chọn thư mục BentoLoop chứa apktool.yml, sau đó bấm Compile. Công cụ đóng gói lại các thành phần cùng libil2cpp.so đã thay và tạo APK để ký, cài đặt.')
para('Quá trình build cần cập nhật Apktool từ 2.6.1 lên 2.12.1 vì bản cũ không nhận thẻ property trong manifest. Với bản mới, bỏ tích Use the aapt2 binary instead of aapt để giao diện không thêm tham số --use-aapt2 đã bị loại bỏ. Khi còn lỗi enableOnBackInvokedCallback not found, đổi tên framework\\1.apk trong bộ nhớ đệm Apktool thành 1.apk.bak rồi Compile lại để công cụ tạo framework mới. Sau các điều chỉnh này, quá trình build và ký hoàn tất.')
retained('bento_report_final', 21, 'APK Easy Tool báo Sign successful sau khi build và ký', 13.8)
para('Mở Compiled APK directory để lấy BentoLoop.apk trong thư mục 2-Recompiled APKs. File này là đầu ra của bước build lại và được chuyển từ máy ảo sang máy thật để cài trên thiết bị Android.')
retained('bento_report_final', 22, 'BentoLoop.apk được tạo trong thư mục 2 Recompiled APKs', 15.4)
heading('5.2 Kiểm thử với chuỗi khác key tìm được', 2)
para('Trên APK mới, nhập vinh123 rồi bấm XÁC NHẬN (VERIFY). Chuỗi này khác BENTO-EASY-2025 nhưng ứng dụng vẫn hiển thị KÍCH HOẠT THÀNH CÔNG! ĐANG VÀO GAME...')
retained('bento_report_final', 23, 'APK đã sửa báo kích hoạt thành công với chuỗi vinh123', 11.3)
para('Kết quả phù hợp với mã VerifyLicense đã sửa để trả về 1 ngay khi được gọi. Ảnh này ghi nhận bản APK mới chấp nhận một chuỗi khác key; ảnh kiểm chứng APK gốc ở mục 3.3 ghi nhận key hợp lệ và màn chơi Level 1.')
table(['Bản ứng dụng', 'Chuỗi thử nghiệm', 'Kết quả ghi nhận'], [
    ('APK gốc', 'BENTO-EASY-2025', 'Kích hoạt thành công và vào Level 1'),
    ('APK đã sửa', 'vinh123', 'Kích hoạt thành công và đang vào game'),
], [4.2, 5.0, 7.7])

heading('6 Nhận xét và phương án phòng tránh')
heading('6.1 Nguyên nhân cơ chế bị vòng tránh', 2)
para('Key được lưu dưới dạng chuỗi trong thư viện phía client và được so sánh trực tiếp trong VerifyLicense. Sau khi xác định nhãn StringLiteral_2770, nội dung BENTO-EASY-2025 có thể đọc từ Listing và đối chiếu với JSON. Các nhánh kiểm tra bổ sung không ngăn được việc tìm key qua nhánh so sánh trực tiếp đang tồn tại.')
para('Quyết định chấp nhận license còn phụ thuộc vào giá trị trả về của một hàm trong mã client. Thay 8 byte đầu hàm đã đưa luồng thực thi đến return 1. APK sau khi build và ký lại vẫn báo kích hoạt thành công với vinh123, cho thấy cơ chế đã quan sát có thể bị vòng tránh bằng cách sửa thư viện.')
heading('6.2 Biện pháp giảm rủi ro', 2)
para('Với ứng dụng có máy chủ, xác minh license theo tài khoản hoặc quyền sử dụng trên máy chủ và tiếp tục kiểm tra quyền tại các thao tác cung cấp dữ liệu, dịch vụ. Quyền truy cập khi đó không chỉ phụ thuộc vào một hàm hoặc một màn hình trong APK.')
para('Loại bỏ key chung và nhánh chấp nhận thử nghiệm khỏi bản phát hành. Nếu cần license offline, có thể dùng license ký số: client giữ khóa công khai để xác minh, còn khóa riêng phát hành license được lưu ngoài ứng dụng. Cách này hạn chế việc nhúng một bí mật chung có thể đọc trực tiếp từ APK, nhưng phần kiểm tra trong client vẫn cần được bảo vệ trước can thiệp.')
para('Kết hợp kiểm tra chữ ký ứng dụng, tính toàn vẹn thư viện và làm rối mã để tăng khả năng phát hiện sửa đổi và tăng chi phí phân tích. Các biện pháp này bổ sung cho cơ chế xác minh quyền, không bảo đảm riêng mã client sẽ ngăn được mọi cách patch.')

p = d.add_paragraph('Tài liệu tham khảo')
p.paragraph_format.first_line_indent = Pt(0)
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.keep_with_next = True
for r in p.runs:
    r.bold = True
    r.font.size = Pt(12)
for ref in [
    'Apktool — Tham số build: https://apktool.org/docs/2.x/cli-parameters/',
    'Apktool — Quản lý framework: https://apktool.org/wiki/in-depth/frameworks/',
    'OWASP MASVS RESILIENCE 2 — Bảo vệ trước can thiệp mã: https://mas.owasp.org/MASVS/controls/MASVS-RESILIENCE-2/',
]:
    p = para(ref)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.keep_together = True
    for r in p.runs:
        r.font.size = Pt(10)

used = {el.get(qn('r:embed')) for el in d.element.xpath('.//a:blip')}
for rid, rel in list(d.part.rels.items()):
    if rel.reltype.endswith('/image') and rid not in used:
        d.part.drop_rel(rid)
update = d.settings.element.find(qn('w:updateFields'))
if update is None:
    update = OxmlElement('w:updateFields')
    d.settings.element.append(update)
update.set(qn('w:val'), 'true')
d.core_properties.title = 'Báo cáo thực hành dịch ngược BentoLoop'
d.core_properties.subject = 'Tìm license key sửa VerifyLicense và kiểm thử APK build lại'
d.save(OUT)
assert hashlib.sha256(REF.read_bytes()).hexdigest() == reference_hash
check = Document(OUT)
text = '\n'.join(check.element.xpath('.//w:t/text()'))
assert not re.search(r'\b(tôi|chúng tôi|chúng ta|mình|bạn)\b', text, re.I)
assert 'SDK' not in text and 'SandSort' not in text
assert len([p for p in check.paragraphs if p.style.name == 'Heading 1']) == 6
assert len(figures) == 17 and len(check.inline_shapes) == 18
hashes = [hashlib.sha256(Path(f['file']).read_bytes()).hexdigest() for f in figures]
assert len(set(hashes)) == len(hashes)
manifest = {'reference': str(REF), 'reference_sha256': reference_hash, 'output': str(OUT), 'figures': figures, 'main_sections': 6, 'sdk_removed': True, 'render_pending': True}
(WORK / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
print('OUTPUT', OUT)
print('FIGURES', len(figures), 'PARAGRAPHS', len(check.paragraphs))
