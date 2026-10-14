from pathlib import Path
from copy import deepcopy
from io import BytesIO
import json, re
from PIL import Image, ImageOps
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent
OUT = ROOT.parent / 'Bao cao dich nguoc SandSort bo sung day du.docx'
SOURCE = Path(r'C:\Users\ADMIN\Downloads\Báo cáo dịch ngược.docx')
TEMP = Path(r'C:\Users\ADMIN\AppData\Local\Temp')
PICDIR = ROOT / 'revision_images'
PICDIR.mkdir(exist_ok=True)
doc = Document(SOURCE)

# Word accepts floating twip values from the source; normalize them to valid OOXML integers.
for s in doc.sections:
    for el in s._sectPr:
        if el.tag in (qn('w:pgMar'), qn('w:pgSz')):
            for key, val in list(el.attrib.items()):
                if re.fullmatch(r'\d+\.\d+', val): el.set(key, str(round(float(val))))

def find(start):
    ps = [p for p in doc.paragraphs if p.text.startswith(start)]
    assert len(ps) == 1, (start, len(ps))
    return ps[0]

def text_set(p, text):
    props = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    run = p.add_run(text)
    if props is not None: run._r.insert(0, props)
    return p

def replace(start, text): return text_set(find(start), text)

def after(anchor, text='', kind='p', page=False):
    el = OxmlElement('w:p'); anchor._p.addnext(el)
    p = Paragraph(el, anchor._parent)
    if kind in ('h2','h3'): p.style = 'Heading 2' if kind == 'h2' else 'Heading 3'
    else: p.style = 'normal'
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12.5 if kind == 'h2' else 12 if kind == 'h3' else 11.5)
    run.font.color.rgb = __import__('docx').shared.RGBColor(0,0,0)
    if kind.startswith('h'): run.bold=True
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = kind.startswith('h')
    p.paragraph_format.widow_control = True
    p.paragraph_format.page_break_before = page
    if kind == 'p': p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def crop(uuid, name, box):
    path = TEMP / ('codex-clipboard-'+uuid+'.png')
    assert path.exists(), path
    im = Image.open(path).convert('RGB').crop(box)
    result=PICDIR/(name+'.png'); im.save(result)
    return result

def composite(uuid, name, entry_box, decomp_box):
    im = Image.open(TEMP/('codex-clipboard-'+uuid+'.png')).convert('RGB')
    entry=im.crop(entry_box); decomp=im.crop(decomp_box)
    w=max(entry.width,decomp.width)
    output=Image.new('RGB',(w,entry.height+18+decomp.height),'white')
    output.paste(entry,(0,0));output.paste(decomp,(0,entry.height+18))
    result=PICDIR/(name+'.png');output.save(result);return result

def stack(uuid,name,boxes):
    im=Image.open(TEMP/('codex-clipboard-'+uuid+'.png')).convert('RGB')
    parts=[im.crop(box) for box in boxes]
    out=Image.new('RGB',(max(p.width for p in parts),sum(p.height for p in parts)+14*(len(parts)-1)),'white')
    y=0
    for part in parts:out.paste(part,(0,y));y+=part.height+14
    dest=PICDIR/(name+'.png');out.save(dest);return dest

pics = {
 'zip':crop('935caa2c-1a4f-4852-87fe-2ee62ee4e593','package_root',(303,140,1320,554)),
 'unity':crop('d7e4c48b-0834-42a4-a6e0-dd8d9f79c770','unity_full_path',(304,140,989,388)),
 'export':crop('b51ea947-06f4-4ca5-8ed3-6e45fb711469','export_png',(307,139,781,348)),
 'binary':crop('91c57b21-2fd7-4266-ae3c-cb682bd0b1ca','binary_path',(308,140,983,481)),
 'metadata':crop('aae747b9-5f5d-4d1a-8ce1-c88f9256c590','metadata_path',(309,141,984,231)),
 'fields':crop('d78c20ac-d0e7-473a-ba7d-7d71e61d617b','uiwingame_click_fields',(627,380,1115,568)),
 'action':crop('ec870b9d-7b62-4b02-bd04-87a773bc9ce5','uiclick_action_field',(615,320,1135,391)),
 'bind_top':composite('1ab1bd3e-23f6-4d7f-8d0c-f2176886c9c9','bind_nextlevel',(641,468,1160,635),(1198,163,1915,753)),
 'bind_x2':crop('7c5f23b6-f18f-4dfe-a815-135d9f7e140b','bind_x2',(1198,163,1915,631)),
 'pointer_meta':crop('d25c0b22-6ad6-42ee-807b-c40c64f2e2db','pointerup_rva',(534,171,1262,282)),
 'pointer':composite('8df438c8-a2ee-44cb-a64a-c57a3976a2f3','pointerup_action',(643,474,1210,607),(1216,164,1912,753)),
 'start':composite('b5f35392-a35f-4c09-87c7-bc7e675a2d69','start_address',(640,266,1273,504),(1286,163,1914,753)),
 'state_entry':crop('19362a5a-7431-4975-817b-e176f0c617d2','statecoin_address',(636,371,1119,752)),
 'wrapper':composite('e1b7564a-2204-4e19-9f79-c5aeb625e77c','nextlevel_wrapper',(641,398,1131,504),(1141,163,1914,753)),
 'claim':composite('17b40228-e4cd-47b4-9d62-90c26e089ea1','trystartclaim_address',(642,381,1273,521),(1284,163,1914,753)),
 'enable':composite('969e21af-abeb-4284-8728-042e37aca839','onenable_address',(642,410,1132,509),(1138,163,1914,601)),
}
pics['start']=stack('b5f35392-a35f-4c09-87c7-bc7e675a2d69','start_focused',[(640,398,1273,504),(1286,163,1914,278),(1286,495,1914,641)])
pics['state_entry']=stack('19362a5a-7431-4975-817b-e176f0c617d2','statecoin_entry_focused',[(636,372,1119,439),(636,684,1119,721)])
pics['pointer']=stack('8df438c8-a2ee-44cb-a64a-c57a3976a2f3','pointerup_focused',[(643,474,1210,577),(1216,164,1912,187),(1216,490,1912,734)])
pics['coin_gain']=crop('d78c20ac-d0e7-473a-ba7d-7d71e61d617b','coin_gain_transform',(291,665,495,715))

def image_run(p, path, width=16.6):
    p.clear();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_after=Pt(4)
    run=p.add_run();run.add_picture(str(path),width=Cm(width))
    for prop in run._r.xpath('.//wp:docPr'):
        prop.set('descr',Path(path).stem.replace('_',' '))
    return p

def figure(anchor, key, caption, width=16.6):
    p=after(anchor);image_run(p,pics[key],width)
    cap=after(p,'Hình NEW. '+caption)
    cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next=False
    cap.paragraph_format.space_after=Pt(8)
    cap.runs[0].italic=True;cap.runs[0].font.size=Pt(10)
    return cap

def replace_figure(number,key,caption,width=16.6):
    cap=find('Hình '+str(number)+'. ')
    prior=Paragraph(cap._p.getprevious(),cap._parent)
    assert prior._p.xpath('.//a:blip'), number
    image_run(prior,pics[key],width)
    text_set(cap,'Hình '+str(number)+'. '+caption)

# Package, Unity input, export evidence and both IL2CPP input paths.
replace('File Explorer hiển thị đầy đủ tên',
 'Gói mẫu được lưu tại C:\\LabGame\\APK_goc\\SandSort_14-9.zip, dung lượng File Explorer hiển thị 110 605 KB. Khi mở gói, các thư mục assets, lib, res, file AndroidManifest.xml và các file classes.dex xuất hiện trực tiếp ở thư mục gốc, phù hợp cấu trúc một gói ứng dụng Android. Ảnh xác nhận tên gói đang lưu là SandSort_14-9.zip; chưa xác nhận tên file APK ban đầu hoặc lịch sử đổi đuôi. Không dùng phần 14-9 trong tên file để suy ra versionName hay versionCode.')
figure(find('Hình 1. '),'zip','Cấu trúc Android hiện trực tiếp trong gói SandSort_14-9.zip')
replace('Bản sao APK được giải nén',
 'Dữ liệu gói ứng dụng được giải nén vào C:\\LabGame\\APK_giai_nen. Các thư mục assets, lib, res và các file classes.dex, classes2.dex cùng AndroidManifest xuất hiện trong thư mục này. assets chứa dữ liệu Unity; lib chứa thư viện native theo kiến trúc. Các file DEX và manifest thuộc phần ứng dụng Android, không thay thế mã IL2CPP của game.')
replace('Trong UABEA, workspace',
 'File Unity được mở để khảo sát bằng UABEA là C:\\LabGame\\APK_giai_nen\\assets\\bin\\Data\\data.unity3d. Workspace data.unity3d chứa các mục sharedassets0.assets, sharedassets1.assets, sharedassets2.assets và các file dữ liệu liên quan. sharedassets0.assets là mục serialized asset bên trong workspace đang mở; không mặc định coi tên này là một file rời ngoài ổ đĩa.')
a=after(find('Hình 8. '),'File Explorer bên dưới xác nhận data.unity3d nằm trong thư mục Data đã giải nén. Ảnh này đối chiếu vị trí trên ổ đĩa; ảnh UABEA ở mục tiếp theo thể hiện workspace, loại asset và Preview.')
figure(a,'unity','Đường dẫn data.unity3d trong thư mục APK_giai_nen')
replace('Sau khi xem Preview, sử dụng',
 'Sau khi xem Preview, sử dụng chức năng Export của UABEA để xuất tài nguyên vào C:\\LabGame\\Asset_xuat. File xuất được xác nhận là bottle_1_Texture2D_sharedassets0.assets.png, có phần mở rộng .png. Thumbnail hình chai tương ứng với asset bottle 1 trong Preview.')
replace_figure(10,'export','File bottle_1_Texture2D_sharedassets0.assets.png trong C:\\LabGame\\Asset_xuat',14.0)
replace('Kết quả gồm vị trí workspace',
 'Kết quả gồm đường dẫn đầy đủ data.unity3d, mục sharedassets0.assets, tên bottle 1, loại Texture2D, ảnh Preview và file PNG đã xuất. Tài nguyên này phục vụ hình ảnh game theo quan sát; việc xác định nơi cộng coin dựa trên metadata và mã native ở các bước sau.')
a=find('File mã native ');text_set(a,'Đầu vào binary: C:\\LabGame\\APK_giai_nen\\lib\\arm64-v8a\\libil2cpp.so. Chọn bản ARM64 trong thư mục arm64-v8a để thống nhất với thư viện được import vào Ghidra.')
figure(a,'binary','Vị trí libil2cpp.so ARM64 dùng làm đầu vào binary',15.8)
a=find('File metadata ');text_set(a,'Đầu vào metadata: C:\\LabGame\\APK_giai_nen\\assets\\bin\\Data\\Managed\\Metadata\\global-metadata.dat.')
figure(a,'metadata','Vị trí global-metadata.dat dùng làm đầu vào metadata',16.6)
replace('Đường dẫn metadata là vị trí',
 'Hai ảnh File Explorer xác nhận tên file và đường dẫn hai đầu vào. Khi chạy Il2CppDumper, chọn libil2cpp.so ở hộp binary, sau đó chọn global-metadata.dat ở hộp metadata. Ảnh vị trí file không phải ảnh hộp chọn của lần chạy dump; đầu ra và kết quả gắn tên trong Ghidra được đối chiếu ở các phần sau.')
after(find('Hình 11. '),
 'Quy trình chuyển từ tên phương thức sang mã máy: tra lớp và chữ ký trong DummyDll bằng dnSpy; tìm cùng phương thức trong dump.cs để lấy RVA và đối chiếu field offset; dùng ghidra.py đọc script.json để gắn nhãn trên libil2cpp.so. Sau đó chọn tên đã gắn trong Symbol Tree hoặc chuyển tới địa chỉ Ghidra = RVA + Image Base, rồi đọc Listing và Decompiler. DummyDll cung cấp cấu trúc, không chứa thân mã native để phân tích hành vi.')

# Address values with matching evidence: retain direct file-offset evidence only where shown.
t=doc.tables[3]
assert t.cell(0,0).text=='Phương thức',t.cell(0,0).text
t.cell(0,2).text='Địa chỉ Ghidra'
for row,val in zip(t.rows[1:],['0x23A0708','0x23A0710','0x23A15D0','0x23A15D4','0x23A16F8']): row.cells[2].text=val
replace('Các phương thức get_Coin, set_Coin, Awake, Start và StateCoin được chọn',
 'Các phương thức get_Coin, set_Coin, Awake, Start và StateCoin được chọn để kiểm tra. get_Coin và set_Coin có ảnh thuộc tính Address và dump.cs bên dưới. Các địa chỉ Ghidra của Awake, Start và StateCoin được đối chiếu bằng ảnh Listing ở mục 6 và 7; RVA tương ứng bằng địa chỉ Ghidra trừ Image Base 0x100000. RVA ở chú thích mã IL của DLL giả và Token metadata không được dùng thay địa chỉ native.')
after(find('Hình 14. '),'Trong ảnh dump.cs, File Offset của get_Coin là 0x229C708 và set_Coin là 0x229C710. Các giá trị này là vị trí trong file ELF, khác địa chỉ Ghidra trong bảng. Những phương thức còn lại được trình bày bằng RVA và địa chỉ Ghidra có ảnh đối chiếu, không bổ sung File Offset khi ảnh chưa thể hiện.')
figure(find('Hình 12. '),'coin_gain','Field _coinGain của UIWinGame có kiểu Transform, thuộc thành phần giao diện',10.0)
replace_figure(20,'start','Start tại 0x23A15D4 đọc GetInt với mặc định 100 và ghi _coin ở +0x20',15.8)
a=after(find('FoodSort.PlayerResourceManager.StateCoin có RVA'),
 'Listing bên dưới xác nhận điểm bắt đầu StateCoin là 0x23A16F8; phần Decompiler tiếp theo thể hiện phép tính và bước lưu số dư.')
figure(a,'state_entry','Điểm bắt đầu PlayerResourceManager.StateCoin tại 0x23A16F8',12.5)
figure(find('UIWinGame.NextLevel tại'),'wrapper','NextLevel tại 0x23B6044 khởi chạy state machine NextLevel d__26',16.0)

# Two adjacent bools are modified by the two-byte stores; callbacks are also cleared.
replace('FoodSort.UIWinGame.TryStartClaim có địa chỉ',
 'FoodSort.UIWinGame.TryStartClaim có địa chỉ Ghidra 0x23B632C, RVA 0x22B632C. Hàm đọc byte ở offset 0xA8 vào cVar1. Nếu giá trị ban đầu bằng 0, phép ghi undefined2 tại +0xA8 bằng 1 đặt _isClaimingReward ở +0xA8 thành true và _isRewardAdCallbackHandled ở +0xA9 thành false. Hai tên field được đối chiếu bằng dnSpy ở mục 7.4. Cuối hàm trả về cVar1 == 0 dựa trên giá trị cờ ban đầu.')
replace_figure(25,'claim','TryStartClaim tại 0x23B632C kiểm tra cờ và xóa callback của các nút',14.0)
a=after(find('Hình 25. '),
 'Sau khi bắt đầu nhận thưởng, các dòng 24–26 ghi 0 vào field +0x20 của đối tượng UIClick ở this +0x50; các dòng 33–36 thực hiện tương tự với đối tượng ở this +0x58 khi nó tồn tại. Đối chiếu field ở mục 7.5 cho thấy đây là thao tác xóa _action của nút tiếp tục và nút nhận thưởng X2. Như vậy, nhánh chuẩn bị vừa đặt cờ vừa bỏ callback đang gắn, giúp hạn chế kích hoạt lặp trong luồng đã quan sát.')
replace('UIWinGame.OnEnable có RVA',
 'UIWinGame.OnEnable có RVA 0x22B5B40, địa chỉ Ghidra 0x23B5B40. Trong nhánh UIManager tồn tại, dòng 71 ghi 0 bằng kiểu undefined2 tại +0xA8, đặt lại cả _isClaimingReward (+0xA8) và _isRewardAdCallbackHandled (+0xA9). Dòng 72 gọi BindClaimActions để gắn lại callback. Không suy ra đây là toàn bộ cơ chế chống nhận lặp hoặc mỗi level chỉ đi qua nhánh này một lần.')
replace_figure(27,'enable','OnEnable tại 0x23B5B40 đặt lại hai cờ và gọi BindClaimActions',16.6)

# Extend the reward trace at its proper location, before the synthesis.
summary=find('7.5. Tổng hợp luồng')
text_set(summary,'7.7. Tổng hợp luồng và giới hạn chứng minh')
summary.paragraph_format.page_break_before=True
anchor=Paragraph(summary._p.getprevious(),summary._parent)
a=after(anchor,'7.5. Gắn callback cho các nút nhận thưởng','h2')
a=after(a,'Trong FoodSort.UIWinGame, dnSpy xác nhận _uIClickNextLevel có FieldOffset 0x50 và _uIClickRewardX2 có FieldOffset 0x58, đều thuộc kiểu UIClick. Đây là các đối tượng được truy cập trong BindClaimActions và TryStartClaim.')
a=figure(a,'fields','Các field UIClick của UIWinGame ở offset 0x50 và 0x58',14.3)
a=after(a,'Trong FoodSort.UIClick, field private Action _action có FieldOffset 0x20. Vì vậy, phép ghi ở địa chỉ đối tượng UIClick +0x20 là cập nhật callback của nút; không phải ghi số dư coin.')
a=figure(a,'action','Field UIClick._action kiểu Action tại offset 0x20',14.3)
a=after(a,'7.5.1. BindClaimActions đăng ký NextLevel và ClaimX2Reward','h3',page=True)
a=after(a,'FoodSort.UIWinGame.BindClaimActions bắt đầu tại địa chỉ Ghidra 0x23B5D30, tương ứng RVA 0x22B5D30 theo Image Base 0x100000. Dòng 22 lấy UIClick ở this +0x50. Dòng 24 tạo System.Action gắn đối tượng UIWinGame hiện tại với phương thức NextLevel; dòng 27 ghi Action vào UIClick +0x20. Đây là đăng ký hàm sẽ chạy khi nút được kích hoạt, chưa phải thực thi NextLevel ngay lúc gắn.')
a=figure(a,'bind_top','BindClaimActions tại 0x23B5D30 gắn NextLevel vào nút _uIClickNextLevel',16.6)
a=after(a,'Với đối tượng ở this +0x58, hàm kiểm tra UnityEngine.Object khác null rồi tạo System.Action gắn với ClaimX2Reward ở dòng 39, sau đó ghi vào field _action ở +0x20. Ảnh chứng minh callback của nhánh X2 được đăng ký; chưa phân tích luồng quảng cáo và lượng thưởng của nhánh đó nên không suy ra nó cộng 200 coin.')
a=figure(a,'bind_x2','BindClaimActions đăng ký ClaimX2Reward cho _uIClickRewardX2',16.6)
a=after(a,'7.6. Sự kiện nhả chuột thực thi callback của UIClick','h2',page=True)
a=after(a,'FoodSort.UIClick.OnPointerUp(PointerEventData) được khai báo trong dnSpy với thuộc tính Address: RVA 0x22B11BC và File Offset 0x22AD1BC. Trong Ghidra hiện tại, địa chỉ bắt đầu tương ứng là 0x23B11BC. RVA nhỏ ở chú thích mã IL của DummyDll không phải địa chỉ native này.')
a=figure(a,'pointer_meta','Thuộc tính Address của UIClick.OnPointerUp trong dnSpy',16.6)
a=after(a,'Ở dòng 48, OnPointerUp gọi RectangleContainsScreenPoint để kiểm tra điểm nhả con trỏ nằm trong vùng hình chữ nhật của nút. Dòng 50 yêu cầu kết quả đúng và _action đọc từ this +0x20 khác null. Dòng 53 thực hiện lời gọi gián tiếp thông qua delegate đó. Đối chiếu với BindClaimActions xác nhận: với đối tượng _uIClickNextLevel, _action đã được gắn với UIWinGame.NextLevel, nên nhả con trỏ hợp lệ sẽ đi vào NextLevel.')
a=figure(a,'pointer','OnPointerUp tại 0x23B11BC kiểm tra vị trí con trỏ và gọi _action',16.6)
a=after(a,'Lời gọi trong OnPointerUp là lời gọi delegate gián tiếp; không cần xuất hiện tên NextLevel trực tiếp tại dòng gọi. Các offset và nơi tạo Action nối được hai phần bằng chứng. UIClick là thành phần dùng chung, nên không kết luận mọi đối tượng UIClick đều gọi NextLevel; kết luận này áp dụng cho nút đã được UIWinGame gắn callback.')

replace('Chuỗi xử lý đã xác định là',
 'Chuỗi xử lý đã xác định là: UIWinGame.OnEnable đặt lại hai cờ và gọi BindClaimActions → gắn NextLevel vào _uIClickNextLevel._action → UIClick.OnPointerUp kiểm tra điểm nhả con trỏ và gọi Action → UIWinGame.NextLevel khởi chạy <NextLevel>d__26.MoveNext → TryStartClaim cho phép lần nhận mới và xóa callback → PlayerResourceManager.StateCoin(true, 100, ...) → cập nhật _coin → PlayerPrefs.SetInt("Coin_SandSort", số dư mới) → PlayerPrefs.Save. Các bước được nối bằng field offset, Action và lời gọi native trong ảnh.')
replace('Chưa xác định đầy đủ sự kiện báo thắng',
 'Đã xác định đối tượng _uIClickNextLevel và đường callback đi tới NextLevel. Phần còn chưa xác định là sự kiện hoặc phương thức xử lý kết quả thắng nào mở UIWinGame, cùng việc kiểm chứng luồng này khi game chạy. Vì vậy, kết luận là tìm được đường nhận thưởng từ nút tiếp tục trong lớp giao diện UIWinGame; chưa khẳng định mọi lần thắng đều tự động cộng ngay 100 coin. Nhánh ClaimX2Reward mới được xác định bước gắn callback, chưa được phân tích đầy đủ.')

t=doc.tables[4]
assert t.cell(0,0).text=='Vị trí',t.cell(0,0).text
for name,addr,meaning in [('BindClaimActions','0x23B5D30','Gắn NextLevel và ClaimX2Reward vào _action của các nút'),('UIClick.OnPointerUp','0x23B11BC','Kiểm tra điểm nhả và gọi delegate _action')]:
    row=t.add_row()
    for cell,value in zip(row.cells,[name,addr,meaning]): cell.text=value
for row in t.rows[1:]:
    if row.cells[0].text=='TryStartClaim': row.cells[2].text='Kiểm tra cờ nhận thưởng và xóa callback trong nhánh bắt đầu'
    if row.cells[0].text=='OnEnable': row.cells[2].text='Đặt lại hai cờ và gắn lại callback trong nhánh đã xem'

replace('Bài thực hành đã nhận diện dữ liệu Unity trong APK',
 'Bài thực hành đã xác định data.unity3d trong thư mục đã giải nén, chọn Texture2D bottle 1 trong sharedassets0.assets, xem Preview và xuất file PNG. Il2CppDumper tạo dữ liệu tra cứu, dnSpy cung cấp cấu trúc lớp và offset, còn Ghidra xác nhận hành vi native. Đối chiếu BindClaimActions, UIClick.OnPointerUp và NextLevel nối được callback của nút tiếp tục với nhánh cộng 100 coin và lưu số dư. Kết luận về kết quả thắng mở giao diện và vai trò máy chủ vẫn cần bằng chứng bổ sung.')
t=doc.tables[6]
assert t.cell(0,0).text=='Nội dung'
for row in t.rows[1:]:
    label=row.cells[0].text
    if label=='Điều kiện nhận thưởng': row.cells[1].text='TryStartClaim kiểm tra _isClaimingReward (+0xA8), đặt hai cờ và xóa callback; OnEnable đặt lại hai cờ và gắn callback trong nhánh đã xem.'
    if label=='Phần chưa xác định': row.cells[1].text='Sự kiện hoàn thành level mở UIWinGame; kiểm chứng động, vai trò server và phần nghiệp vụ của nhánh X2.'
row=t.add_row();row.cells[0].text='Callback của nút tiếp tục';row.cells[1].text='BindClaimActions gắn NextLevel cho _uIClickNextLevel; OnPointerUp gọi _action khi điểm nhả hợp lệ và callback khác null.'

# Renumber figures in body order and apply consistent paragraph/table layout.
figure_count=0
for p in doc.paragraphs:
    if re.match(r'^Hình (?:\d+|NEW)\. ',p.text):
        figure_count+=1
        text_set(p,re.sub(r'^Hình (?:\d+|NEW)\.',f'Hình {figure_count}.',p.text))
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=False
        p.paragraph_format.page_break_before=False
        for r in p.runs: r.italic=True;r.font.size=Pt(10)
    if p.style.name.startswith('Heading'):
        p.paragraph_format.keep_with_next=True
        for r in p.runs:r.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
for t in doc.tables[1:]:
    borders=t._tbl.tblPr.find(qn('w:tblBorders'))
    if borders is None:borders=OxmlElement('w:tblBorders');t._tbl.tblPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=borders.find(qn('w:'+edge))
        if el is None:el=OxmlElement('w:'+edge);borders.append(el)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:el.set(qn('w:'+k),v)
    for row in t.rows:
        pr=row._tr.get_or_add_trPr()
        for height in list(pr.findall(qn('w:trHeight'))): pr.remove(height)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.widow_control=True
                for r in p.runs:
                    r.font.name='Times New Roman';r.font.size=Pt(10.5)
    header=t.rows[0]._tr.get_or_add_trPr()
    if header.find(qn('w:tblHeader')) is None: header.append(OxmlElement('w:tblHeader'))

# Replace the source's manual page-break-only paragraphs with controlled section breaks.
# This prevents blank pages after larger replacement figures.
for i,p in enumerate(list(doc.paragraphs)):
    if i>3 and not p.text.strip() and not p._p.xpath('.//a:blip') and p._p.xpath('.//w:br[@w:type="page"]'):
        p._p.getparent().remove(p._p)
intro=find('Bài thực hành khảo sát cấu trúc APK Unity')
intro.paragraph_format.page_break_before=True
replace('Kết quả phân tích tĩnh cho thấy',
 'Kết quả phân tích tĩnh cho thấy PlayerResourceManager.Start nạp coin từ PlayerPrefs với khóa Coin_SandSort, dùng 100 khi chưa tồn tại khóa. BindClaimActions gắn NextLevel vào callback của nút tiếp tục và UIClick.OnPointerUp thực thi callback khi nhả con trỏ hợp lệ. Phần MoveNext của NextLevel gọi StateCoin với state = true và coin = 100; StateCoin cập nhật _coin rồi lưu số dư mới. Một Texture2D hình chai đã được xuất thành file PNG. Chưa xác định sự kiện kết quả thắng mở UIWinGame hoặc xác nhận hành vi khi game chạy.')
new_pages=('1.1 ','1.2 ','1.3 ','1.4 ','2.2.','2.2 ','2.3.','2.3 ','Đầu ra và cách','4.2.','5.3 ','6.2.','6.3.','7.2.','7.3.','7.4.','7.5.','7.6.','7.7.')
for p in doc.paragraphs:
    if p.style.name=='Heading 1':
        p.paragraph_format.page_break_before=not p.text.startswith('1.')
    elif p.style.name in ('Heading 2','Heading 3'):
        p.paragraph_format.page_break_before=p.text.startswith(new_pages)
    if p.text.strip() and not p.text.startswith('Hình '):
        following=p._p.getnext()
        if following is not None and following.xpath('.//a:blip'):
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.keep_together=True
    if '\n' in p.text or p.text.startswith(('Địa chỉ Ghidra =','_coin =','int change =','# -*-')):
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent=Cm(0)
    if p.text.startswith(('GNU Build ID','Địa chỉ Ghidra =','int change =','# -*-')):
        for r in p.runs:r.font.name='Consolas';r.font.size=Pt(9)
    if p.style.name not in ('Heading 1','Heading 2','Heading 3') and p.text.strip() and not p.text.startswith('Hình ') and '\n' not in p.text:
        # Keep prose consistent with the supplied report's sans-serif body.
        for r in p.runs:r.font.name='Arial';r.font.size=Pt(11)

# Keep the short rejected-candidates table together instead of stranding its final row.
for row in doc.tables[5].rows[:-1]:
    for cell in row.cells:
        for p in cell.paragraphs:p.paragraph_format.keep_with_next=True

# Move the StateCoin signature next to its actual implementation instead of a sparse overflow page.
p=find('StateCoin có chữ ký public void')
signature=p.text
p._p.getparent().remove(p._p)
p=find('FoodSort.PlayerResourceManager.StateCoin có RVA')
text_set(p,'FoodSort.PlayerResourceManager.StateCoin(bool state, int coin, string placement = "unknown", string groupPlacement = "unknown") có RVA 0x22A16F8, địa chỉ Ghidra 0x23A16F8. param_1 là đối tượng; param_2 là bool state; param_3 là int coin. Hai tham số được lưu vào đối tượng phụ trước khi sử dụng; số dư vẫn nằm ở param_1 +0x20.')
# The directory overview needs less space than the Preview and native-code figures.
cap=find('Hình 9. Cấu trúc thư mục APK_giai_nen')
p=Paragraph(cap._p.getprevious(),cap._parent)
for extent in p._p.xpath('.//wp:extent'):
    cx=int(extent.get('cx'));cy=int(extent.get('cy'));new=int(Cm(14.0))
    extent.set('cx',str(new));extent.set('cy',str(round(cy*new/cx)))
for extent in p._p.xpath('.//a:xfrm/a:ext'):
    cx=int(extent.get('cx'));cy=int(extent.get('cy'));new=int(Cm(14.0))
    extent.set('cx',str(new));extent.set('cy',str(round(cy*new/cx)))

settings=doc.settings.element
update=settings.find(qn('w:updateFields'))
if update is None: update=OxmlElement('w:updateFields');settings.append(update)
update.set(qn('w:val'),'true')
doc.core_properties.title='Báo cáo dịch ngược SandSort'
doc.core_properties.subject='Khảo sát asset Unity và logic coin IL2CPP ARM64'
doc.save(OUT)

check=Document(OUT)
text='\n'.join(p.text for p in check.paragraphs)
assert 'Chưa xác định đầy đủ sự kiện báo thắng hoặc nút cụ thể' not in text
assert 'phần mở rộng ảnh không hiện' not in text
assert len([p for p in check.paragraphs if p.text.startswith('Hình ')])==figure_count
assert all(x in text for x in ['0x23B5D30','0x23B11BC','0x23B6044','0x23A16F8','bottle_1_Texture2D_sharedassets0.assets.png'])
(ROOT/'revision_summary.json').write_text(json.dumps({'output':str(OUT),'figures':figure_count,'tables':len(check.tables),'images':len(check.inline_shapes),'new_images':{k:str(v) for k,v in pics.items()}},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(OUT),'figures':figure_count,'tables':len(check.tables),'images':len(check.inline_shapes)},ensure_ascii=False))
