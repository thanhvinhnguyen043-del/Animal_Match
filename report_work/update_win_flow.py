from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import re, json, shutil, hashlib
from PIL import Image
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent
WORK = ROOT / 'win_flow_update'
WORK.mkdir(exist_ok=True)
SOURCE = Path(r'C:\Users\ADMIN\Downloads\Báo cáo dịch ngược.docx')
TARGET = ROOT.parent / 'Báo cáo dịch ngược SandSort cập nhật luồng thắng.docx'
shutil.copy2(SOURCE, WORK / 'source_backup.docx')
source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
d = Document(SOURCE)
original = list(d.paragraphs)
normal_props = deepcopy(original[190]._p.pPr)
heading_props = deepcopy(original[189]._p.pPr)

def replace(p, text):
    if p.runs:
        p.runs[0].text = text
        for run in p.runs[1:]:
            run.text = ''
    else:
        p.add_run(text)

replace(original[132], 'StateCoin thực hiện việc cộng hoặc trừ số dư theo tham số truyền vào. Nhánh nhận thưởng thông thường trên UIWinGame gọi hàm này với state = true và coin = 100. Phần phân tích dưới đây đối chiếu cả thao tác nhận thưởng và luồng kiểm tra thắng dẫn đến việc mở giao diện này.')
replace(original[189], '7.8. Tổng hợp luồng và giới hạn chứng minh')
replace(original[190], 'Khi _isWin chưa được đặt và CheckWin trả về true, TryWinLevel chuyển sang CompleteLevelAsWin. Phần MoveNext của luồng này kiểm tra _isWin và _isLose, đánh dấu thắng, chờ 500 ms rồi bật GameObject của _uIWinGame nếu đối tượng tồn tại. Khi giao diện được kích hoạt, OnEnable có nhánh đặt lại hai cờ nhận thưởng và gọi BindClaimActions nếu UIManager tồn tại.')
replace(original[191], 'Sau khi callback đã được gắn, thao tác nhả con trỏ hợp lệ trên nút tiếp tục làm OnPointerUp gọi NextLevel. NextLevel khởi chạy state machine; phần MoveNext kiểm tra TryStartClaim và gọi StateCoin với state = true, coin = 100 khi được phép nhận thưởng. StateCoin cộng 100 vào _coin, lưu số dư mới bằng PlayerPrefs.SetInt với khóa Coin_SandSort rồi gọi PlayerPrefs.Save.')
replace(original[192], 'Luồng đã được nối từ kiểm tra thắng, mở UIWinGame đến nhận và lưu thưởng qua nút tiếp tục. Lượng thưởng 100 thuộc nhánh này, không phải thao tác tự động cộng ngay khi CheckWin trả về true. Số coin hiển thị và dữ liệu lưu khi game chạy chưa được kiểm chứng. CheckWin còn có nhiều nhánh và trường dữ liệu chưa được đối chiếu đầy đủ; nhánh X2 mới được xác định ở bước gắn callback. Nội dung hai chuỗi mô tả nguồn thưởng và việc đặt lại trạng thái ở các luồng khác cũng chưa được phân tích hết.')
replace(original[200], 'Dữ liệu PlayerPrefs khi game chạy và số coin hiển thị sau khi nhận thưởng chưa được kiểm tra. Các nhánh thưởng khác, nội dung hai chuỗi mô tả nguồn thưởng và các hàm thống kê trong StateCoin cũng chưa được phân tích đầy đủ. Luồng kiểm tra thắng đã được nối với giao diện nhận thưởng, nhưng chưa đủ để kết luận mọi lần thắng đều tự động cộng ngay 100 coin hoặc có thể can thiệp thành công vào số dư.')
replace(original[205], 'Qua bài thực hành, đã thực hiện được các bước từ giải nén gói ứng dụng, tìm và xuất asset Unity đến phân tích mã IL2CPP. UABEA được dùng để xem trước và xuất texture; Il2CppDumper tạo dữ liệu tra cứu; dnSpy giúp xác định lớp, phương thức và field; Ghidra được dùng để đọc mã native và lần theo tham chiếu. Kết hợp các công cụ này có thể tìm nơi lưu một giá trị, đối chiếu field offset với lệnh đọc ghi và lần theo luồng xử lý từ giao diện đến phần cập nhật dữ liệu.')
replace(original[206], 'Với SandSort, đã xác định được cách nạp số coin khi vào game, giá trị mặc định 100 khi chưa có dữ liệu lưu và cách cộng trừ, lưu số dư. Luồng kiểm tra thắng cũng đã được nối đến thao tác mở UIWinGame và nhận 100 coin qua nút tiếp tục. Kết quả cho thấy cần đọc cả MoveNext của phương thức bất đồng bộ và đối chiếu tên field để hiểu đúng mã giả. Phân tích hiện tại dựa trên mã tĩnh; số coin thực tế khi chạy game, các nhánh thưởng khác và cơ chế xác thực ngoài phần đã xem vẫn cần được kiểm tra thêm.')

anchor = original[189]
def insert(text='', heading=False):
    el = OxmlElement('w:p')
    el.append(deepcopy(heading_props if heading else normal_props))
    anchor._p.addprevious(el)
    p = Paragraph(el, anchor._parent)
    if heading:
        p.style = d.styles['Heading 2']
        p.paragraph_format.page_break_before = False
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    if text:
        p.add_run(text)
    return p

temp = Path(r'C:\Users\ADMIN\AppData\Local\Temp')
manifest = []
def figure(file, box, number, caption, width=17):
    src = temp / file
    im = Image.open(src)
    crop = im.crop(box)
    destination = WORK / f'figure_{number}.png'
    crop.save(destination)
    p = insert()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Pt(0)
    p.paragraph_format.right_indent = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(destination), width=Cm(width))
    picture._inline.docPr.set('descr', caption)
    cp = insert(f'Hình {number}. {caption}')
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.first_line_indent = Pt(0)
    cp.paragraph_format.left_indent = Pt(0)
    cp.paragraph_format.right_indent = Pt(0)
    cp.paragraph_format.line_spacing = 1
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(6)
    cp.paragraph_format.keep_together = True
    cp.paragraph_format.keep_with_next = False
    for r in cp.runs:
        r.bold = True
        r.italic = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0,0,0)
    manifest.append({'number':number,'source':str(src),'crop':box,'caption':caption})

insert('7.7. Điều kiện thắng và mở giao diện nhận thưởng', heading=True)
insert('FoodSort.LevelManager.TryWinLevel bắt đầu tại địa chỉ Ghidra 0x239B76C. Hàm kiểm tra byte tại this +0xF2 đang bằng 0 rồi gọi CheckWin. Khi CheckWin trả về true, lệnh b tại 0x239B78C chuyển đến CompleteLevelAsWin ở 0x239B2CC. Đây là tham chiếu mã được tìm từ danh sách References của CompleteLevelAsWin; các dòng tham chiếu dữ liệu trong danh sách không phải lời gọi thực thi.')
figure('codex-clipboard-d914f47f-e3a5-4cf8-abf8-67a17183c154.png', (1136,165,1917,630), 41, 'TryWinLevel kiểm tra cờ thắng và kết quả CheckWin trước khi xử lý thắng')
insert('CheckWin bắt đầu tại 0x239BDC4 và chỉ xét tiếp khi byte tại this +0xF3 bằng 0. Hàm có nhánh theo IsNoHolesNormalLevel, duyệt danh sách BottleController và gọi IsEmpty. Những đoạn khác kiểm tra Header.CheckAnyBottle, IsAnyBottleFlying và CheckCompleteAllBottles. Đoạn cuối còn so sánh giá trị tại this +0x104 với số phần tử của danh sách tại this +0x148. Vì các trường này và thân các hàm con chưa được đối chiếu hết, điều kiện thắng chưa được rút gọn thành một quy tắc duy nhất. Riêng dòng 119 trong Hình 42 kiểm tra kết quả CheckCompleteAllBottles bằng 0, nên không thể chỉ dựa vào tên hàm để kết luận tất cả chai hoàn thành thì thắng.')
figure('codex-clipboard-0002c7fe-3c10-42f4-873f-e532ed3bcb5f.png', (1164,449,1914,751), 42, 'Đoạn cuối CheckWin kiểm tra các hàm con và phép so sánh trước khi trả kết quả')
insert('Trong dnSpy, LevelManager có field bool _isWin ở offset 0xF2 và bool _isLose ở offset 0xF3. Hai khai báo này xác nhận tên của các cờ được đọc trong TryWinLevel, CheckWin và CompleteLevelAsWin.MoveNext. Field int _levelDisplay ở offset 0xF4 cũng xuất hiện trong cùng ảnh; phép tăng giá trị tại offset này trong luồng thắng là cập nhật giá trị màn chơi, không phải tăng coin.')
figure('codex-clipboard-9b351ec8-54ab-4b58-8a44-63af89b78ae6.png', (520,145,1110,692), 43, 'Các field _isWin, _isLose và _levelDisplay của LevelManager trong dnSpy', width=12)
insert('CompleteLevelAsWin ở 0x239B2CC khởi chạy state machine <CompleteLevelAsWin>d__205. Phần xử lý nằm trong MoveNext tại 0x239E110. Để xác định đối tượng được bật trong hàm này, field tại LevelManager +0x68 được đối chiếu trong dnSpy: đó là private UIWinGame _uIWinGame, có thuộc tính SerializeField.')
figure('codex-clipboard-5044e0d6-6cb9-4f9a-8b78-5994c9eeebe7.png', (520,145,1200,652), 44, 'Field _uIWinGame thuộc kiểu UIWinGame tại offset 0x68 của LevelManager', width=14)
insert('Ở nhánh bắt đầu của MoveNext, nếu _isWin hoặc _isLose đã là true thì hàm chuyển đến phần kết thúc. Nếu cả hai cờ đều false, hàm đặt _isWin thành true và gọi Task.Delay với giá trị 500 ms. Điều kiện này giúp tránh bắt đầu lại phần xử lý thắng khi màn chơi đã có trạng thái kết thúc.')
figure('codex-clipboard-14b42320-578e-4b10-a943-dd725055f51a.png', (1113,480,1916,752), 45, 'MoveNext kiểm tra hai cờ, đánh dấu thắng và chờ 500 ms')
insert('Sau khi chờ, dòng 81–87 lấy đối tượng tại LevelManager +0x68, gọi Component.get_gameObject rồi GameObject.SetActive với giá trị true nếu đối tượng tồn tại. Kết hợp với Hình 44, đoạn này xác nhận CompleteLevelAsWin bật GameObject của UIWinGame. Đoạn mã mở giao diện chưa trực tiếp cộng coin; thao tác cộng 100 nằm trong nhánh NextLevel.MoveNext đã phân tích ở mục 7.2.')
figure('codex-clipboard-74e81434-25b0-4793-a465-f2eb24887538.png', (1113,459,1916,751), 46, 'MoveNext lấy GameObject của _uIWinGame tại offset 0x68 và gọi SetActive true')

# Extend the existing address summary without changing the rest of the report.
table = d.tables[4]
rows = [
    ('TryWinLevel', '0x239B76C', 'Kiểm tra chưa thắng và gọi CheckWin; chuyển đến CompleteLevelAsWin khi điều kiện đúng'),
    ('CheckWin', '0x239BDC4', 'Trả kết quả kiểm tra thắng theo các nhánh đã khảo sát'),
    ('CompleteLevelAsWin', '0x239B2CC', 'Khởi chạy state machine xử lý thắng'),
    ('CompleteLevelAsWin.MoveNext', '0x239E110', 'Đánh dấu thắng, chờ 500 ms và bật GameObject của UIWinGame'),
]
for values in reversed(rows):
    row = deepcopy(table.rows[1]._tr)
    table.rows[0]._tr.addnext(row)
    for cell, text in zip(table.rows[1].cells, values):
        p = cell.paragraphs[0]
        replace(p, text)
        for extra in list(cell.paragraphs[1:]):
            extra._p.getparent().remove(extra._p)

# Keep the previous formatting requests throughout the whole document.
for style_name in ('Heading 1', 'Heading 2', 'Heading 3'):
    d.styles[style_name].paragraph_format.page_break_before = False
for el in d.element.xpath('.//w:p'):
    p = Paragraph(el, d)
    if p.style.name.startswith('Heading') and p.text.strip():
        p.paragraph_format.page_break_before = False
    is_image = bool(el.xpath('.//w:drawing | .//w:pict'))
    is_caption = bool(re.match(r'^Hình\s+\d+\.', p.text))
    if is_image or is_caption:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Pt(0)
        p.paragraph_format.right_indent = Pt(0)
    if is_image:
        p.paragraph_format.keep_with_next = True
    if is_caption:
        p.paragraph_format.keep_together = True
        for r in p.runs:
            r.bold = True
            r.italic = True
            r.font.color.rgb = RGBColor(0,0,0)

d.save(TARGET)
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == source_hash
with ZipFile(SOURCE) as before, ZipFile(TARGET) as after:
    media = [n for n in before.namelist() if n.startswith('word/media/')]
    assert all(before.read(n) == after.read(n) for n in media), 'Original media must be preserved'
check = Document(TARGET)
captions = [p.text for p in check.paragraphs if re.match(r'^Hình\s+\d+\.', p.text)]
assert len(captions) == 46
assert len(check.inline_shapes) == len(Document(SOURCE).inline_shapes) + 6
assert not any('chưa xác định hàm xử lý kết quả thắng nào mở UIWinGame' in p.text or 'chưa xác định được hàm mở màn hình thắng' in p.text for p in check.paragraphs)
(WORK / 'manifest.json').write_text(json.dumps({'source':str(SOURCE),'source_sha256':source_hash,'output':str(TARGET),'new_figures':manifest},ensure_ascii=False,indent=2),encoding='utf-8')
print('Updated report:', TARGET)
print('Figures:', len(captions), '; all original image files preserved.')
